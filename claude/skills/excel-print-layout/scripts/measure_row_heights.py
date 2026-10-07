# -*- coding: utf-8 -*-
# 实测行高：找出「合并单元格 + wrap 多行文本」但行高不足的行，用 Excel 自己
# 算权威自然行高。
#
# 为什么不用 AutoFit / .Text / 字符估算：
#   - 对合并单元格所在行 AutoFit 会被重置为默认行高，无效；
#   - Cell.Text 不反映 wrap 截断（行高再小仍返回全文）；
#   - 字符宽度估算对多行文本会系统性高估 1-2 行。
# 唯一可信的方法是「临时单元格 AutoFit」：在空行上按合并宽度 + 原字体排入文本，
# 让 Excel 算自然行高。实测值会原样落到 JSON，供 apply_row_heights.py 写回。
import sys, io, argparse, json, time, math, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl, win32com.client


def char_w(c):
    return 2 if ord(c) > 0xFF else 1


def text_lines(text, width):
    # 仅用于候选筛选（判断是否多行），不影响最终高度（最终由 Excel 实测）
    if not text:
        return 0
    n = 0
    for seg in str(text).split("\n"):
        if not seg:
            n += 1
            continue
        n += max(1, (sum(char_w(c) for c in seg) + width - 1) // width)
    return n


def parse_args():
    ap = argparse.ArgumentParser(description="实测 Excel 合并单元格 wrap 行所需自然行高（需本机 Excel）。")
    ap.add_argument("input", help="输入 xlsx 路径")
    ap.add_argument("--output", help="实测结果 JSON 路径；缺省写 input 同目录 temp/row_heights_measured.json")
    ap.add_argument("--hpt-base", type=float, default=13.2,
                    help="每行高度基数（pt），用于候选筛选近似，不影响最终实测值。默认 13.2")
    ap.add_argument("--tolerance", type=float, default=2.0,
                    help="实测高度超出现行高该差值（pt）才记为需调整。默认 2.0")
    return ap.parse_args()


def build_candidates(wb, hpt_base):
    # {sheet_title: {str(row): (col, text, fname, size, merged_width, cur_height)}}
    cand = {}
    for ws in wb.worksheets:
        colw = {}
        for c in range(1, ws.max_column + 1):
            cd = ws.column_dimensions.get(openpyxl.utils.get_column_letter(c))
            colw[c] = cd.width if (cd and cd.width) else 8.43
        merged = list(ws.merged_cells.ranges)

        def merged_width(r, c):
            for m in merged:
                if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col:
                    return sum(colw[x] for x in range(m.min_col, m.max_col + 1))
            return None

        rows = {}
        for r in range(1, ws.max_row + 1):
            rd = ws.row_dimensions.get(r)
            cur = rd.height if (rd and rd.height) else 13.8
            best = None  # (need, col, text, fname, size, width, cur)
            for c in range(1, ws.max_column + 1):
                cc = ws.cell(row=r, column=c)
                v = cc.value
                if not isinstance(v, str) or not v.strip():
                    continue
                mw = merged_width(r, c)
                wflag = bool(cc.has_style and cc.alignment.wrap_text)
                if not (wflag or mw):
                    continue
                w = mw if mw else colw[c]
                if text_lines(v, max(1, int(w))) <= 1:
                    continue
                sz = cc.font.size if (cc.has_style and cc.font and cc.font.size) else 10
                fname = cc.font.name if (cc.has_style and cc.font and cc.font.name) else "Times New Roman"
                need = round(text_lines(v, max(1, int(w))) * hpt_base * sz / 10.0, 1)
                if best is None or need > best[0]:
                    best = (need, c, v, fname, sz, round(w, 1), round(cur, 1))
            if best:
                rows[str(r)] = best[1:]  # (col, text, fname, size, width, cur)
        if rows:
            cand[ws.title] = rows
    return cand


def main():
    args = parse_args()
    out = args.output or os.path.join(
        os.path.dirname(os.path.abspath(args.input)), "temp", "row_heights_measured.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)

    wb = openpyxl.load_workbook(args.input, data_only=True)
    cand = build_candidates(wb, args.hpt_base)
    wb.close()
    print("候选实测行数:", sum(len(v) for v in cand.values()), flush=True)

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    result = {}
    TR = 500  # 临时行，用于 AutoFit 实测
    t0 = time.time()
    try:
        wbk = excel.Workbooks.Open(os.path.abspath(args.input), ReadOnly=True, UpdateLinks=0)
        for si, (title, rows) in enumerate(cand.items(), 1):
            ws = wbk.Worksheets(title)
            sres = {}
            for r, (c, text, fname, size, wdt, cur) in rows.items():
                ws.Columns(1).ColumnWidth = max(1, wdt)
                tc = ws.Cells(TR, 1)
                tc.Value = text
                tc.Font.Name = fname
                tc.Font.Size = size
                tc.WrapText = True
                ws.Rows(TR).AutoFit()
                real = round(ws.Rows(TR).RowHeight, 1)
                tc.Value = None
                if real > cur + args.tolerance:
                    sres[int(r)] = round(max(real, 13.8), 1)
            result[title] = sres
            done = sum(len(v) for v in result.values())
            print("[%d/%d] %-34s 需调整 %3d 行  累计 %4d  %.0fs" % (
                si, len(cand), title, len(sres), done, time.time() - t0), flush=True)
        wbk.Close(False)
    finally:
        excel.Quit()

    json.dump(result, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("完成，共 %d 行，耗时 %.1f 分钟" % (sum(len(v) for v in result.values()), (time.time() - t0) / 60))
    print("输出:", out)


if __name__ == "__main__":
    main()
