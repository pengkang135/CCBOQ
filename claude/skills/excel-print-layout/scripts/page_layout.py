# -*- coding: utf-8 -*-
# Excel 打印排版：设置打印宽度（fitToWidth）+ 插入竖向分页符。
#
# 为什么纯 ZIP+lxml 而不是 openpyxl 保存：openpyxl 保存会丢失媒体文件、
# printerSettings、webextensions、metadata、calcChain、sharedStrings，导致
# Excel 无法打开该文件。openpyxl 在这里只用于读值判定锚点行，写操作全走
# 原生 XML 重写，其余 zip 条目原样搬运。
import sys, io, argparse, re, zipfile, shutil, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from lxml import etree
import openpyxl
from openpyxl.utils import column_index_from_string

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def parse_cols(spec):
    # "A,B" -> [1, 2]（1-based 列号）
    return [column_index_from_string(x.strip().upper()) for x in spec.split(",") if x.strip()]


def parse_args():
    ap = argparse.ArgumentParser(
        description="Excel 打印排版：页宽 + 分页符（纯 XML 写，openpyxl 只读值）。")
    ap.add_argument("input", help="输入 xlsx 路径")
    ap.add_argument("--output", help="输出路径；缺省则先备份 input 到同目录 temp/ 后就地覆盖")
    ap.add_argument("--fit-width", type=int, default=None,
                    help="fitToWidth 值（1 = 内容缩放到一页宽）。不传则不动打印宽度")
    ap.add_argument("--break-after", action="append", default=[],
                    help="正则：每个匹配行的【之后】插分页符（该行留在本页底）。可多次")
    ap.add_argument("--break-before", action="append", default=[],
                    help="正则：每个匹配行的【之前】插分页符（该行成为新页首行）。可多次")
    ap.add_argument("--rows", action="append", default=[],
                    help="显式行号（1-based，逗号分隔），在每个行号的【之后】插分页符")
    ap.add_argument("--cols", default="A,B",
                    help="正则匹配时检查的列（逗号分隔，如 'A,B'），默认 A,B")
    return ap.parse_args()


def row_matches(ws, r, patterns, cols):
    # 行 r 的 cols 列中任一单元格文本命中任一正则即返回 True
    texts = []
    for c in cols:
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str):
            texts.append(v)
    for t in texts:
        for pat in patterns:
            if re.search(pat, t):
                return True
    return False


def collect_breaks(args, wb):
    # 返回 {sheet_title: [分页符 brk id ...]}。brk id = 分页前最后一行号。
    # --break-after 匹配行 r -> brk id = r；--break-before 匹配行 r -> brk id = r-1。
    cols = parse_cols(args.cols)
    pat_after = [re.compile(p) for p in args.break_after]
    pat_before = [re.compile(p) for p in args.break_before]
    rows_explicit = []
    for spec in args.rows:
        rows_explicit += [int(x) for x in spec.split(",") if x.strip()]

    result = {}
    for ws in wb.worksheets:
        ids = set()
        for r in range(1, ws.max_row + 1):
            if pat_after and row_matches(ws, r, pat_after, cols):
                ids.add(r)
            if pat_before and row_matches(ws, r, pat_before, cols):
                ids.add(r - 1)
        for r in rows_explicit:
            ids.add(r)
        ids.discard(0)
        result[ws.title] = sorted(ids)
    return result


def ensure_page_setup_fit(root, fit_width):
    # pageSetup 设 fitToWidth/fitToHeight=0、去 scale；sheetPr/pageSetUpPr 设 fitToPage=1
    ps = root.find("{%s}pageSetup" % NS)
    if ps is None:
        ps = etree.SubElement(root, "{%s}pageSetup" % NS)
    ps.set("fitToWidth", str(fit_width))
    ps.set("fitToHeight", "0")
    if "scale" in ps.attrib:
        del ps.attrib["scale"]
    spr = root.find("{%s}sheetPr" % NS)
    if spr is None:
        spr = etree.Element("{%s}sheetPr" % NS)
        root.insert(0, spr)
    psup = spr.find("{%s}pageSetUpPr" % NS)
    if psup is None:
        psup = etree.Element("{%s}pageSetUpPr" % NS)
        spr.append(psup)
    psup.set("fitToPage", "1")


def write_row_breaks(root, ids):
    rbk = root.find("{%s}rowBreaks" % NS)
    if rbk is None:
        rbk = etree.Element("{%s}rowBreaks" % NS)
        anchor = root.find("{%s}headerFooter" % NS)
        if anchor is None:
            anchor = root.find("{%s}pageSetup" % NS)
        if anchor is None:
            anchor = root.find("{%s}pageMargins" % NS)
        if anchor is not None:
            anchor.addnext(rbk)
        else:
            root.append(rbk)
    for br in list(rbk):
        rbk.remove(br)
    rbk.set("count", str(len(ids)))
    rbk.set("manualBreakCount", str(len(ids)))
    for r in ids:
        b = etree.SubElement(rbk, "{%s}brk" % NS)
        b.set("id", str(r))
        b.set("max", "16383")
        b.set("man", "1")


def main():
    args = parse_args()
    src = args.input
    out = args.output or src
    if args.output is None:
        d = os.path.dirname(os.path.abspath(src))
        bak = os.path.join(d, "temp", os.path.splitext(os.path.basename(src))[0] + "_页宽分页前备份.xlsx")
        os.makedirs(os.path.dirname(bak), exist_ok=True)
        shutil.copy2(src, bak)
        print("备份:", bak)

    # 内存模式（勿用 read_only=True：随机 cell() 访问是 O(n)，20 个 sheet 会卡死）
    wb = openpyxl.load_workbook(src, data_only=True)
    breaks = collect_breaks(args, wb)
    titles = [ws.title for ws in wb.worksheets]
    wb.close()

    zin = zipfile.ZipFile(src)
    items = []
    for n in zin.namelist():
        data = zin.read(n)
        m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", n)
        if m:
            idx = int(m.group(1))
            if idx <= len(titles):
                title = titles[idx - 1]
                ids = breaks.get(title, [])
                fit = args.fit_width
                if ids or fit is not None:
                    root = etree.fromstring(data)
                    if fit is not None:
                        ensure_page_setup_fit(root, fit)
                    if ids:
                        write_row_breaks(root, ids)
                    data = etree.tostring(root, xml_declaration=True,
                                          encoding="UTF-8", standalone=True)
        items.append((n, data, zin.getinfo(n).compress_type))
    zin.close()

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data, ct in items:
            zout.writestr(n, data)

    for title in titles:
        ids = breaks.get(title, [])
        if ids or args.fit_width is not None:
            print("%-36s 分页符 %d 个 %s" % (title, len(ids), str(ids)[:60]))
    print("输出:", out)


if __name__ == "__main__":
    main()
