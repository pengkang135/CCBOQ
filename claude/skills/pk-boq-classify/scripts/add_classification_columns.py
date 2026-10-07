#!/usr/bin/env python3
"""
Add 16 classification columns to a BOQ Excel file (SortKey integrated).
Dynamically detects last data column, inserts classification columns with formulas and styles.
"""
import argparse
import re
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    ("BQ Code",     "formula"),
    ("Dept1",       "static"),
    ("Dept2",       "static"),
    ("Dept3",       "static"),
    ("Discipline",  "ai"),
    ("SortKey",     "ai"),
    ("Category",    "ai"),
    ("Subcategory", "ai"),
    ("Element",     "ai"),
    ("Material",    "ai"),
    ("Spec",        "ai"),
    ("Description", "formula"),
    ("Unit",        "ref"),
    ("Quantity",    "ref"),
    ("Rate",        "ref"),
    ("Amount",      "ref"),
]

COL_WIDTHS = {
    "BQ Code": 36, "Dept1": 28, "Dept2": 28, "Dept3": 28,
    "Discipline": 14, "SortKey": 10, "Category": 14, "Subcategory": 16,
    "Element": 22, "Material": 22, "Spec": 22,
    "Description": 48, "Unit": 8, "Quantity": 10, "Rate": 12, "Amount": 14,
}

FILL_BLUE = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FILL_RED  = PatternFill(start_color="C0504D", end_color="C0504D", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
FONT_HEADER = Font(name="Microsoft YaHei UI", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Microsoft YaHei UI", size=9)
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_DATA = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

COL_DESC = 3    # C: Item Description
COL_UNIT = 5    # E: Unit
COL_QTY  = 6    # F: Quantity
COL_RATE = 8    # H: Unit Rate (inc duties and taxes on imported goods)
COL_AMOUNT = 16 # P: Total Price


def find_last_data_col(ws):
    last_col = 0
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value is not None:
            last_col = col
    return last_col


def make_header_styles():
    styles = {}
    for hdr, typ in HEADERS:
        fill = FILL_BLUE if typ in ("formula", "static", "ref") else FILL_RED
        styles[hdr] = {
            "font": FONT_HEADER,
            "fill": fill,
            "alignment": ALIGN_HEADER,
            "border": THIN_BORDER,
        }
    return styles


def add_classification_columns(in_path, out_path=None, sheet_name=None):
    if out_path is None:
        out_path = in_path

    wb = openpyxl.load_workbook(in_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    print(f"Sheet: {ws.title}, rows: {ws.max_row}, cols: {ws.max_column}")

    last_col = find_last_data_col(ws)
    print(f"Last data column: {get_column_letter(last_col)} (col {last_col})")

    start_col = last_col + 2
    print(f"Classification columns start at: {get_column_letter(start_col)}")

    header_styles = make_header_styles()

    for i, (hdr, typ) in enumerate(HEADERS):
        col = start_col + i
        cell = ws.cell(row=1, column=col, value=hdr)
        sty = header_styles[hdr]
        cell.font = sty["font"]
        cell.fill = sty["fill"]
        cell.alignment = sty["alignment"]
        cell.border = sty["border"]
        ws.column_dimensions[get_column_letter(col)].width = COL_WIDTHS.get(hdr, 12)

    ws.row_dimensions[1].height = 28

    parent_l1 = ""
    parent_l2 = ""
    parent_l3 = ""

    desc_col_letter = get_column_letter(COL_DESC)
    unit_col_letter = get_column_letter(COL_UNIT)

    formula_rows = 0
    data_rows = 0

    for row in range(2, ws.max_row + 1):
        desc_val = ws.cell(row=row, column=COL_DESC).value
        desc = str(desc_val) if desc_val else ""
        unit_val = ws.cell(row=row, column=COL_UNIT).value
        qty_val = ws.cell(row=row, column=COL_QTY).value

        if desc.startswith("【") and "】" in desc:
            parent_l1 = desc.replace("【", "").replace("】", "").strip()
            parent_l2 = ""
            parent_l3 = ""
        if desc.startswith("《") and "》" in desc:
            parent_l2 = desc.replace("《", "").replace("》", "").strip()
            parent_l3 = ""
        if desc.startswith("{") and "}" in desc:
            parent_l3 = desc.replace("{", "").replace("}", "").strip()

        has_unit = unit_val is not None and str(unit_val).strip() != ""
        has_qty = False
        if qty_val is not None:
            try:
                has_qty = float(qty_val) > 0
            except (ValueError, TypeError):
                pass
        is_heading = desc.startswith(("【", "《", "{"))
        is_valid = has_unit and has_qty and not is_heading

        if is_valid:
            data_rows += 1

            # BQ Code: simple Excel formula (no regex, no _xlfn)
            bq_formula = f'=TRIM({desc_col_letter}{row})&"|"&LOWER(TRIM({unit_col_letter}{row}))'
            ws.cell(row=row, column=start_col, value=bq_formula)

            # Dept1-3 (static, from parent tracking)
            ws.cell(row=row, column=start_col + 1, value=parent_l1)
            ws.cell(row=row, column=start_col + 2, value=parent_l2)
            ws.cell(row=row, column=start_col + 3, value=parent_l3)

            # Description: strip leading numbering + trailing periods (regex req'd, compute in Python)
            desc_str = str(desc_val or "").strip()
            desc_clean = re.sub(r'^\s*(?:\d+\.\d+(?:\.\d+)*|-)\s+', '', desc_str)
            desc_clean = re.sub(r'[。.]+$', '', desc_clean).strip()
            ws.cell(row=row, column=start_col + 11, value=desc_clean)

            # Ref columns: static values
            ws.cell(row=row, column=start_col + 12, value=str(unit_val).strip() if unit_val else "")
            ws.cell(row=row, column=start_col + 13, value=qty_val)
            ws.cell(row=row, column=start_col + 14,
                    value=ws.cell(row=row, column=COL_RATE).value)
            ws.cell(row=row, column=start_col + 15,
                    value=ws.cell(row=row, column=COL_AMOUNT).value)

            formula_rows += 1

            # AI columns: yellow fill hint (offsets 4-10: Discipline, SortKey, Category..Spec)
            for ai_offset in range(4, 11):
                cell = ws.cell(row=row, column=start_col + ai_offset)
                cell.fill = FILL_YELLOW

        # Apply data font and border to all 16 classification cells
        for offset in range(16):
            cell = ws.cell(row=row, column=start_col + offset)
            cell.font = FONT_DATA
            cell.border = THIN_BORDER

    wb.save(out_path)
    print(f"Done: {data_rows} data rows, {formula_rows} formula rows written")
    print(f"Output: {out_path}")
    return out_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Add BOQ classification columns")
    parser.add_argument("input", help="Input .xlsx file")
    parser.add_argument("--output", "-o", help="Output .xlsx (default: overwrite input)")
    parser.add_argument("--sheet", help="Sheet name (default: active sheet)")
    args = parser.parse_args()
    add_classification_columns(args.input, args.output, args.sheet)
