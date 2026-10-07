"""DEPRECATED as of 2026-08-08.

SortKey insertion and Discipline split are now integrated into:
- add_classification_columns.py (16-col layout includes SortKey from the start)
- classify_boq_engine.py (_finalize_item() handles SortKey + Discipline split)

This script is kept for reference only. Do NOT use in new pipelines.
"""

import argparse
import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "pk-boq" / "scripts"))

import openpyxl
from openpyxl_utils import clean_save

DISC_TO_BOOK = {
    "Civil & Decoration": "A",
    "Civil": "A",
    "Decoration": "A",
    "MEP Installation": "B",
    "External & Municipal": "C",
    "Marine & Waterway": "D",
    "Preliminaries & General": "PRELIM",
    "Preliminaries": "PRELIM",
}

BOOK_PATTERNS = {
    "A": "企业定额_A册_*.sqlite",
    "B": "企业定额_B册_*.sqlite",
    "C": "企业定额_C册_*.sqlite",
    "D": "企业定额_D册_*.sqlite",
}


def find_header_row(ws):
    for r in range(1, min(ws.max_row + 1, 10)):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val and str(val).strip() == "Discipline":
                return r
    return None


def find_col_by_header(ws, header_row, name):
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        if val and str(val).strip() == name:
            return c
    return None


def build_mappings(db_dir):
    div_zone = {}
    cat_to_code = {}

    for book, pattern in BOOK_PATTERNS.items():
        matches = list(Path(db_dir).glob(pattern))
        if not matches:
            continue
        conn = sqlite3.connect(str(matches[0]))
        rows = conn.execute(
            "SELECT code, name_EN FROM division WHERE name_EN IS NOT NULL ORDER BY code"
        ).fetchall()
        conn.close()

        for code, name_en in rows:
            cat_to_code[(book, name_en)] = code
            if book == "A":
                div_num = int(code.split(".")[1])
                zone = "Decoration" if 20 <= div_num <= 29 else "Civil"
                div_zone[(book, name_en)] = zone

    return div_zone, cat_to_code


def main():
    parser = argparse.ArgumentParser(description="Post-process classified BOQ: split Discipline + insert SortKey")
    parser.add_argument("input", help="Classified BOQ xlsx file")
    parser.add_argument("--output", "-o", help="Output file path (default: input with _sortkey suffix)")
    parser.add_argument("--sheet", default="合并报表", help="Sheet name (default: 合并报表)")
    parser.add_argument("--db-dir", default=r"E:\Code\Norms-AI\db",
                        help="Norms DB directory")
    args = parser.parse_args()

    div_zone, cat_to_code = build_mappings(args.db_dir)
    print(f"Division zones: {len(div_zone)}, Category→Code: {len(cat_to_code)}")

    wb = openpyxl.load_workbook(args.input)
    ws = wb[args.sheet]

    header_row = find_header_row(ws)
    if not header_row:
        print("ERROR: Could not find header row with 'Discipline'")
        return

    col_disc = find_col_by_header(ws, header_row, "Discipline")
    col_cat = find_col_by_header(ws, header_row, "Category")
    if not col_disc or not col_cat:
        print(f"ERROR: Could not find Discipline (col {col_disc}) or Category (col {col_cat})")
        return

    # Insert SortKey before Category
    ws.insert_cols(col_cat)
    ws.cell(row=header_row, column=col_cat, value="SortKey")

    disc_updated = 0
    sortkey_written = 0

    for r in range(header_row + 1, ws.max_row + 1):
        disc = ws.cell(row=r, column=col_disc).value
        cat = ws.cell(row=r, column=col_cat + 1).value  # shifted after insert

        if not cat:
            continue

        cat_str = str(cat).strip()
        disc_str = str(disc).strip() if disc else ""

        book = DISC_TO_BOOK.get(disc_str, "")

        # Split Civil & Decoration
        if disc_str == "Civil & Decoration" and book == "A":
            zone = div_zone.get((book, cat_str))
            if zone:
                ws.cell(row=r, column=col_disc, value=zone)
                disc_updated += 1

        # Write SortKey
        if book == "PRELIM":
            sk = "PRELIM"
        else:
            sk = cat_to_code.get((book, cat_str), "")
        if sk:
            ws.cell(row=r, column=col_cat, value=sk)
            sortkey_written += 1

    output = args.output or str(Path(args.input).with_stem(Path(args.input).stem + "_sortkey"))
    clean_save(wb, output)
    wb.close()

    print(f"Discipline split: {disc_updated}, SortKey written: {sortkey_written}")
    print(f"Saved: {output}")


if __name__ == "__main__":
    main()
