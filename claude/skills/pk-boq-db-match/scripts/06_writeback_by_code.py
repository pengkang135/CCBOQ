"""V2 Writeback aligned by BQ Code (F column) — resilient against row shifts.

Rules:
    - Copy 套价版.xlsx → 套价版_AI版.xlsx
    - For each leaf row in current source, read F (BQ Code), look up in V2 result JSON
    - Skip if N (人工套价) is a POSITIVE number (user's real price); N=0/empty/None → write AI
    - Write O-W (O 置信度, P 名称, Q 单位, R 单价 THB, S 币种, T 项目, U 供应商, V 日期, W 理由)
    - Never touch A-N
    - Clear O-W for rows without a V2 match (fresh state)
"""
import json, shutil, math
from pathlib import Path
import openpyxl

ROOT = Path(r"f:\BaiduSyncdisk\5.报价\2026-7 泰国CMI移动MOD2数据中心")
SRC = ROOT / "4 BQ" / "2026-7-29 第二版" / "2026-07-29_CMI_MOD2_BOQ_合并报表_套价版.xlsx"
DST = ROOT / "4 BQ" / "2026-7-29 第二版" / "2026-07-29_CMI_MOD2_BOQ_合并报表_套价版_AI版.xlsx"
REPORT_DIR = ROOT / "4 BQ" / "2026-7-29 第二版" / "2026-07-29_CMI_MOD2_BOQ_合并报表_套价版_套价报告"

def load_v2():
    jsons = sorted(REPORT_DIR.glob("V2_results_by_bqcode_*.json"))
    if not jsons:
        raise FileNotFoundError("no V2 by-bqcode JSON")
    return json.loads(jsons[-1].read_text(encoding="utf-8"))["by_bq_code"]

def positive_num(v):
    if v is None:
        return False
    if isinstance(v, str):
        s = v.strip()
        if not s or s.startswith("="):
            return False
        try:
            return float(s) > 0
        except ValueError:
            return False
    if isinstance(v, (int, float)):
        return not math.isnan(v) and v > 0
    return False

def main():
    v2 = load_v2()
    print(f"Loaded {len(v2)} V2 records (BQ Code keyed)")

    shutil.copy2(SRC, DST)
    print(f"Copied: {DST.name}")

    wb = openpyxl.load_workbook(DST, keep_vba=False)
    ws = wb["UniqueBQ"]

    written = 0
    skipped_manual = 0
    no_v2_data = 0
    cleared = 0

    # Iterate current source rows (post-user-edits)
    for r in range(5, ws.max_row + 1):
        A = ws.cell(r, 1).value      # discipline
        if A == "总计":
            # Clear grand total row's O-W
            for col in range(15, 24):
                ws.cell(r, col).value = None
            ws.cell(r, 15).value = "skip"
            ws.cell(r, 23).value = "总计行 grand total"
            continue

        G = ws.cell(r, 7).value       # qty
        F = ws.cell(r, 6).value       # BQ Code
        N = ws.cell(r, 14).value      # manual price

        # Non-leaf rows (no qty)
        if not (isinstance(G, (int, float)) and G > 0):
            continue

        # User's real price → protect
        if positive_num(N):
            skipped_manual += 1
            continue

        # Clear O-W first (fresh state for any leaf row)
        for col in range(15, 24):
            ws.cell(r, col).value = None
        cleared += 1

        # Look up in V2 by BQ Code
        code = str(F).strip() if F else ""
        rec = v2.get(code)
        if not rec:
            no_v2_data += 1
            continue

        status = rec.get("status", "no_match")
        ws.cell(r, 15).value = status  # O 置信度
        price = rec.get("eff_price_thb") or rec.get("matched_price_thb")

        if status in ("high", "medium", "low", "estimated"):
            ws.cell(r, 16).value = rec.get("matched_name")           # P 名称
            ws.cell(r, 17).value = rec.get("matched_unit")           # Q 单位
            if price:
                ws.cell(r, 18).value = round(float(price), 4)        # R 单价
            ws.cell(r, 19).value = "THB"                             # S 币种
            ws.cell(r, 20).value = rec.get("matched_project")        # T 项目
            ws.cell(r, 21).value = rec.get("matched_supplier")       # U 供应商
            ws.cell(r, 22).value = rec.get("matched_date")           # V 日期
        ws.cell(r, 23).value = (rec.get("reasoning") or "")[:500]    # W 理由
        written += 1

    wb.save(DST)
    print(f"\n=== Writeback stats ===")
    print(f"Written (O-W filled with AI):     {written}")
    print(f"Skipped (N filled with real price): {skipped_manual}")
    print(f"Cleared no V2 data (BQ code mismatch): {no_v2_data}")
    print(f"Total O-W cleared: {cleared}")
    print(f"Output: {DST.name}")

if __name__ == "__main__":
    main()
