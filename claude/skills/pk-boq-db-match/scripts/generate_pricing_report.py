"""
Generic AI pricing report HTML generator.

Usage:
    python generate_pricing_report.py --config config.json

Config JSON structure:
{
    "project_name": "CMI MOD2",
    "root_dir": "f:/path/to/project/4 BQ/2026-7-29 第二版",
    "src_xlsx": "CMI_MOD2_BOQ.xlsx",
    "out_xlsx_name": "CMI_MOD2_BOQ_AI套价.xlsx",
    "results_json": "套价报告/temp/v18_merged_results_refreshed.json",
    "boq_items_json": "套价报告/temp/boq_items.json",
    "sheet_name": "UniqueBQ",
    "header_row": 5,
    "currency": "THB",
    "currency_label": "百万THB",
    "discipline_map": {
        "Civil": "土建工程",
        "Finishes": "装饰工程",
        "External Works": "室外工程",
        "MEP": "安装工程",
        "Preliminary": "开办费"
    },
    "discipline_order": ["土建工程", "装饰工程", "室外工程", "安装工程", "开办费"],
    "data_source_desc": "CostSpread MongoDB (31,004条泰国 THB)",
    "writeback_desc": "Q-Z 列",
    "subtitle_lines": [
        "项目名称从 MongoDB projectName 字段提取 | 4层流水线匹配 + 人工修正"
    ],
    "version_note": "本版更新：项目名称从 MongoDB projectName 字段提取...",
    "rules_text": "土方回填/余土外运优先场内倒运（83.46 THB/m3）",
    "mermaid": {
        "boq_label": "UniqueBQ 309项",
        "data_source_label": "MongoDB 报价库<br/>31,004条 泰国 THB",
        "rules_label": "人工修正规则<br/>土方场内倒运 · 单位换算"
    },
    "human_rules_extra": ""
}
"""

import json, html, re, sys, argparse
from datetime import datetime
from pathlib import Path
from collections import Counter, defaultdict

import openpyxl

# ═══════════════════════════════════════════════════════════════
#  Reasoning / keyword cleaners (shared across all projects)
# ═══════════════════════════════════════════════════════════════

def _extract_keywords(raw):
    parts = re.split(r'[;|,]', raw)
    keywords = []
    skip = re.compile(
        r'^(?:u_exact|u_mismatch|x-lang|base_match|g_\w+|s_\d+\.?\d*|price_ok|no_change|'
        r'cand_\w+|rejected\s*\d*|ROLLBACK|pre-stage7|stage7|no_viable|'
        r'[①-⑨]|no_match\s*[→><]|sim=[\d.]+|total=[\d.]+|'
        r'unit_gap|unit_incompatible|unit_mismatch|no thickness|no viable|no-candidate|'
        r'\[FIX\]|\[MongoDB|\[no-candidate\]|\[known-gap\]|confirm|stage\d|\s*)$'
    )
    for p in parts:
        p = p.strip()
        if not p or skip.search(p):
            continue
        p = re.sub(r'no_match\s*[→><]\s*\w+', '', p).strip()
        if len(p) >= 2:
            keywords.append(p)
    return keywords


def _build_reasoning(status, boq_desc, matched_name="", conv_note="", matched_project="", matched_supplier=""):
    name_short = (matched_name or "")[:60]
    proj = (matched_project or "")
    supp = (matched_supplier or "")
    desc_short = (boq_desc or "")[:80]
    source = ""
    if proj and proj not in ("Unknown", "未知"):
        source = f"，来源: {proj}"
        if supp and not supp.startswith("{"):
            source += f"({supp})"

    if status == "no_match":
        return f"价格库中未找到匹配{desc_short}的条目，建议人工询价或专业分包报价"
    if status == "known_gap":
        return f"价格库缺少此类目，需专业分包报价或历史项目数据补充"
    if status == "unit_gap":
        return f"匹配到类似条目但单位维度不同无法换算，需人工判断"
    if conv_note:
        return f"匹配{name_short}，{conv_note}{source}"
    labels = {
        "high": f"精确匹配{name_short}{source}",
        "medium": f"匹配{name_short}，品类一致{source}",
        "low": f"近似匹配{name_short}，建议核实材质规格{source}",
        "estimated": f"基于{name_short}估算换算，建议核实参数{source}",
    }
    return labels.get(status, "")


def _clean_reasoning(raw, boq_desc, status, matched_name="", conv_note="", matched_project="", matched_supplier=""):
    if not raw:
        return _build_reasoning(status, boq_desc, matched_name, conv_note, matched_project, matched_supplier)

    has_markers = re.search(
        r'[\|⑥⑦⑧⑨]|sim=[\d.]+|total=[\d.]+|ROLLBACK|u_exact|u_mismatch|'
        r'x-lang|base_match|no_viable|pre-stage|rejected\s*\d|unit_gap|unit_incompatible', raw
    )
    if not has_markers and len(raw) > 30:
        return raw[:300]

    m = re.search(r'\[MongoDB补搜\]\s*([^|⑥⑦⑧]+)', raw)
    mongo_info = m.group(1).strip().rstrip(';| ') if m else ""

    auditor_notes = [
        m.group(1).strip() for m in re.finditer(r'⑦auditor:\s*([^|⑥]+)', raw)
        if m.group(1).strip() not in ('confirm_haiku', 'confirm', 'pass')
    ]

    cleaned = raw
    for pat in [
        r'\bsim=[\d.]+\b', r'\btotal=[\d.]+\b', r'\bs_\d+\.?\d*\b', r'\bg_\w+\b',
        r'\bu_\w+\b', r'\bprice_ok\b', r'\bcand_\w+\b', r'\bno\s+viable[^|;]*',
        r'\(\d+\s*(?:filtered|rejected|候选)?\)', r'\|?\s*⑤单位换算[^|;]*\|?',
        r'\|?\s*\[MongoDB补搜\]\s*', r'\|?\s*\[no-candidate\][^|;]*\|?',
        r'\|?\s*\[known-gap\][^|;]*\|?', r'\|?\s*\[FIX\][^|;]*\|?',
        r'⑥retry\s*\w*\s*\([^)]*\)', r'⑥\w+:\s*[^|;]*', r'⑦\w+:\s*[^|;]*',
        r'ROLLBACK:\s*[^|;]*', r'\bpre-stage7\b', r'\|?\s*stage7审核\s*\|?',
        r'\b\w+:\s*no_match\s*[→><]\s*\w+', r'\brejected\s*\d*\b',
        r'(?<!\w)(?:u_exact|u_mismatch|x-lang|base_match|g_unknown|price_ok|'
        r'no_change|unit_incompatible|incompatible|no_viable)\s*[;|]?',
        r'unit\s*(?:mismatch|incompatible|_gap)[^|;]*',
        r'\b[mklt][g23]?\s*[↔→←⟷⟶]\s*[mklt][g23]?\s*(?:via|with|out)?\s*\d*\s*\w*\b',
        r'm<>count|m<>m3|m↔m3|dimension incompatible',
    ]:
        cleaned = re.sub(pat, '', cleaned)

    cleaned = re.sub(r'\s*\|\s*', '; ', cleaned)
    cleaned = re.sub(r';\s*;', ';', cleaned)
    cleaned = re.sub(r'^[;,\s|]+', '', cleaned)
    cleaned = re.sub(r'[;,\s|]+$', '', cleaned)
    cleaned = re.sub(r'\s{2,}', ' ', cleaned)
    cleaned = re.sub(r'\(\s*\)', '', cleaned)
    cleaned = re.sub(r'[，,]\s*[。.]', '。', cleaned)
    cleaned = re.sub(r'[。.]\s*[，,]', '，', cleaned)
    cleaned = re.sub(r'[。.]{2,}', '。', cleaned)
    cleaned = re.sub(r'[，,]{2,}', '，', cleaned)
    cleaned = re.sub(r';\s*[。.]', '。', cleaned)
    cleaned = cleaned.strip(' \t\n\r;|')

    keywords = _extract_keywords(cleaned)
    parts = []
    if mongo_info and len(mongo_info) > 5:
        parts.append(mongo_info)
    if keywords:
        kw_str = '、'.join(keywords[:4])
        if len(kw_str) > 3:
            parts.append(f'匹配关键词: {kw_str}')
    for note in auditor_notes:
        if len(note) > 3:
            parts.append(note)

    if not parts or sum(len(p) for p in parts) < 25:
        return _build_reasoning(status, boq_desc, matched_name, conv_note, matched_project, matched_supplier)

    result = '; '.join(parts)
    result = re.sub(r'^(中置信|低置信|高置信)度匹配[，。]?\s*', '', result)
    return result[:300]


# ═══════════════════════════════════════════════════════════════
#  Unit normalization + conversion notes
# ═══════════════════════════════════════════════════════════════

def norm_unit(u):
    """Normalize unit strings for comparison."""
    if not u:
        return ""
    u = u.strip().lower()
    u = u.replace(".", "").replace(" ", "")
    u = u.replace("²", "2").replace("³", "3")
    cn_map = dict.fromkeys("个套座根只件台组支块把盏副对", "no")
    cn_map.update({"套": "set", "组": "set", "副": "set", "对": "set"})
    for cn, en in cn_map.items():
        u = u.replace(cn, en)
    u = u.replace("㎡", "m2").replace("㎥", "m3")
    alias = {
        "sqm": "m2", "sq.m": "m2", "squaremeter": "m2",
        "cum": "m3", "cu.m": "m3", "cubicmeter": "m3",
        "l": "litre", "liter": "litre",
        "kg": "kg", "t": "t", "ton": "t", "tonne": "t",
        "m": "m", "m2": "m2", "m3": "m3",
        "no": "no", "nr": "no", "ea": "no", "each": "no",
        "set": "set", "lot": "set", "ls": "set",
    }
    return alias.get(u, u)


def _extract_thickness_mm(desc_text, matched_text):
    for src in [desc_text, matched_text]:
        m = re.search(r'(\d+)\s*mm\s*(?:Thick|thick|Thk|thk)', src)
        if m:
            return float(m.group(1))
    return None


def _extract_height_mm(desc_text):
    m = re.search(r'(\d+)\s*mm\s*(?:High|high|H\.?[=:]?\s*\d+)', desc_text)
    if not m:
        m = re.search(r'(\d+)\s*mm\s*High', desc_text)
    if m:
        try:
            return float(m.group(1))
        except (ValueError, IndexError):
            pass
    return None


def build_conv_note(mu, bu, op, cp, mu_raw, bu_raw, desc_text, matched_text):
    """Build human-readable conversion note for unit/price transformations."""
    if mu == bu:
        if cp and op and abs(cp - op) > 0.01:
            ratio = cp / op
            if abs(ratio - round(ratio)) < 0.02 and ratio > 1.01:
                layers = round(ratio)
                return f"{op:,.2f} THB/{bu_raw} x {layers}层 = {cp:,.2f} THB/{bu_raw}"
            if abs(ratio - 1 / round(1 / ratio)) < 0.02 and ratio < 0.99:
                div = round(1 / ratio)
                return f"{op:,.2f} THB/{bu_raw} / {div} = {cp:,.2f} THB/{bu_raw}"
            if 4.5 < ratio < 5.5:
                return f"CNY {op:,.2f} x ~5.0 = {cp:,.2f} THB"
            if 32 < ratio < 35:
                return f"USD {op:,.2f} x ~33 = {cp:,.2f} THB"
            if ratio > 20:
                return f"USD {op:,.2f} x ~{ratio:.0f} = {cp:,.2f} THB"
            return f"{op:,.2f} THB/{bu_raw} -> {cp:,.2f} THB/{bu_raw} (系数 {ratio:.2f}x)"
        return ""

    thickness = _extract_thickness_mm(desc_text, matched_text)
    height_mm = _extract_height_mm(desc_text)

    if mu == "t" and bu == "kg":
        return f"{op:,.2f} THB/t / 1000 = {op / 1000:,.2f} THB/kg"
    if mu == "kg" and bu == "t":
        return f"{op:,.2f} THB/kg x 1000 = {op * 1000:,.2f} THB/t"
    if mu == "m3" and bu == "m2" and thickness:
        t = thickness / 1000
        return f"{op:,.2f} THB/m3 x {t:.3f}m({thickness:.0f}mm厚) = {op * t:,.2f} THB/m2"
    if mu == "m2" and bu == "m3" and thickness:
        t = thickness / 1000
        return f"{op:,.2f} THB/m2 / {t:.3f}m({thickness:.0f}mm厚) = {op / t:,.2f} THB/m3"
    if mu == "m3" and bu == "t":
        return f"{op:,.2f} THB/m3 / 2.4t/m3 = {op / 2.4:,.2f} THB/t"
    if mu == "t" and bu == "m3":
        return f"{op:,.2f} THB/t x 2.4t/m3 = {op * 2.4:,.2f} THB/m3"
    if mu in ("l", "litre") and bu == "m2":
        return f"{op:,.2f} THB/L -- 涂料/防水材料，需按实际涂布率(m2/L)换算为m2单价"
    if mu == "m2" and bu in ("l", "litre"):
        return f"{op:,.2f} THB/m2 -- 需按实际涂布率(m2/L)反算为L单价"
    if mu == "m2" and bu in ("no", "each", "set"):
        return f"{op:,.2f} THB/m2 -- 需按单件面积换算为{bu}单价"
    if mu == "m" and bu == "m2" and height_mm:
        h = height_mm / 1000
        return f"{op:,.2f} THB/m / {h:.3f}m({height_mm:.0f}mm高) = {op / h:,.2f} THB/m2"
    if mu == "m2" and bu == "m" and height_mm:
        h = height_mm / 1000
        return f"{op:,.2f} THB/m2 x {h:.3f}m({height_mm:.0f}mm高) = {op * h:,.2f} THB/m"

    note = f"{op:,.2f} THB/{mu_raw} -> 目标单位{bu_raw}"
    if thickness:
        note += f" (含{thickness:.0f}mm厚度参数)"
    return note


# ═══════════════════════════════════════════════════════════════
#  HTML/CSS template  (the polished V19 design)
# ═══════════════════════════════════════════════════════════════

CSS = """\
*{box-sizing:border-box;margin:0;padding:0;}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI','Microsoft YaHei','PingFang SC',sans-serif;margin:0;padding:24px;background:#f0f2f5;color:#1e293b;line-height:1.6;}
.container{max-width:2400px;margin:0 auto;}
.header{background:linear-gradient(135deg,#0f172a 0%,#1e3a5f 50%,#312e81 100%);color:#fff;padding:32px 40px;border-radius:12px;margin-bottom:24px;box-shadow:0 4px 24px rgba(15,23,42,0.25);}
.header h1{margin:0 0 12px 0;font-size:24px;font-weight:700;letter-spacing:0.5px;}
.header p{margin:4px 0;opacity:0.85;font-size:13px;line-height:1.7;}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:16px;margin-bottom:24px;}
.card{background:#fff;border-radius:10px;padding:20px 16px;text-align:center;box-shadow:0 1px 4px rgba(0,0,0,0.06);border:1px solid #e8ecf1;transition:transform .15s ease,box-shadow .15s ease;}
.card:hover{transform:translateY(-2px);box-shadow:0 6px 20px rgba(0,0,0,0.1);}
.card .num{font-size:28px;font-weight:700;margin:4px 0;line-height:1.2;}
.card .label{font-size:13px;color:#64748b;margin-top:6px;font-weight:500;}
.card:nth-child(1) .num{color:#2563eb;}
.card:nth-child(2) .num{color:#059669;}
.card:nth-child(3) .num{color:#7c3aed;}
.card:nth-child(4) .num{color:#dc2626;}
.card:nth-child(5) .num{color:#d97706;}
.card:nth-child(6) .num{color:#0891b2;}
.section{background:#fff;border-radius:10px;padding:24px 28px;margin-bottom:24px;box-shadow:0 1px 4px rgba(0,0,0,0.06);border:1px solid #e8ecf1;}
.section h2{margin:0 0 16px 0;font-size:17px;font-weight:600;color:#1e3a5f;border-bottom:3px solid #2563eb;padding-bottom:10px;}
.section p{font-size:14px;line-height:1.85;margin:8px 0;}
table{width:100%;border-collapse:collapse;font-size:12px;}
.summary-table th.num{text-align:center;padding-right:12px;}
.summary-table td{padding:6px 8px;}
th{background:#1e293b;color:#f1f5f9;padding:10px 8px;text-align:left;font-weight:600;font-size:12px;position:sticky;top:0;z-index:100;white-space:nowrap;}
td{padding:7px 8px;border-bottom:1px solid #f1f5f9;vertical-align:middle;}
tr:nth-child(even) td{background:#f8fafc;}
tr:hover td{background:#eef2ff!important;}
.elem{max-width:110px;font-size:11px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;}
.desc{max-width:280px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;line-height:1.4;}
.num{text-align:center;font-variant-numeric:tabular-nums;white-space:nowrap;line-height:1.4;}
.conv{max-width:130px;font-size:10px;color:#7c3aed;word-break:break-word;line-height:1.45;}
.reason{max-width:240px;font-size:11px;color:#64748b;line-height:1.5;word-break:break-word;}
.price{text-align:center;font-variant-numeric:tabular-nums;white-space:nowrap;font-size:12px;font-weight:500;color:#0f172a;}
.unit-cell{text-align:center;white-space:nowrap;}
.badge{font-size:11px;padding:3px 10px;border-radius:4px;font-weight:600;color:#fff;text-align:center;display:inline-block;white-space:nowrap;letter-spacing:0.3px;}
.b-high{background:#059669;} .b-medium{background:#d97706;} .b-low{background:#dc2626;}
.b-estimated{background:#7c3aed;} .b-no_match{background:#6b7280;} .b-skip{background:#9ca3af;}
tr[data-status="no_match"] td{background:#fef2f2;}
tr[data-status="low"] td{background:#fffbeb;}
tr[data-status="estimated"] td{background:#faf5ff;}
tr[data-status="known_gap"] td{background:#f8fafc;}
tr[data-status="unit_gap"] td{background:#fff7ed;}
.alert{background:#fef2f2;border:1px solid #fecaca;border-left:4px solid #dc2626;color:#991b1b;padding:12px 16px;border-radius:8px;margin:10px 0;font-size:13px;line-height:1.6;}
.note{background:#eff6ff;border:1px solid #bfdbfe;border-left:4px solid #3b82f6;color:#1e40af;padding:12px 16px;border-radius:8px;margin:10px 0;font-size:13px;line-height:1.6;}
.detail-table{border-radius:8px;border:1px solid #e2e8f0;overflow:clip;}
.summary-table th,.summary-table td{text-align:center;}
.summary-table th:first-child,.summary-table td:first-child{text-align:left;}
.cat-header td{background:#e0e7ff!important;color:#312e81;font-size:12px;font-weight:600;padding:10px 8px;border-bottom:2px solid #6366f1;}
.footer{text-align:center;color:#94a3b8;font-size:12px;padding:20px;}
.mermaid-wrap{background:#f8fafc;border-radius:10px;padding:20px 12px 8px 12px;margin:16px 0;overflow-x:auto;}
.mermaid-wrap .mermaid{display:flex;justify-content:center;min-width:900px;}"""

MERMAID_JS = '<script src="https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.min.js"></script>'
MERMAID_INIT = '<script>mermaid.initialize({startOnLoad:true,theme:\'default\',securityLevel:\'loose\'});</script>'

STATUS_LABEL = {
    "high": "高置信", "medium": "中置信", "low": "低置信",
    "estimated": "估值换算", "no_match": "未匹配",
    "known_gap": "已知无覆盖", "unit_gap": "单位不匹配", "skip": "跳过",
}
BADGE_CLASS = {
    "high": "b-high", "medium": "b-medium", "low": "b-low",
    "estimated": "b-estimated", "no_match": "b-no_match",
    "known_gap": "b-no_match", "unit_gap": "b-no_match", "skip": "b-skip",
}
DETAIL_HEADER_COLS = [
    ("行", None), ("分项", None), ("BOQ描述", None), ("单位", None), ("数量", None),
    ("置信度", None), ("匹配名称", None), ("匹配单位", None), ("除税单价", None),
    ("币种", None), ("换算说明", None), ("报价时间", None), ("供应商", None),
    ("来源项目", None), ("匹配说明", None),
]
DETAIL_HEADER_COLS_GAP = [
    ("行", None), ("分项", None), ("BOQ描述", None), ("单位", None), ("数量", None),
    ("状态", None), ("匹配名称", None), ("匹配单位", None), ("除税单价", None),
    ("币种", None), ("换算说明", None), ("报价时间", None), ("供应商", None),
    ("来源项目", None), ("匹配说明", None),
]
SUMMARY_HEADER_COLS = [
    ("专业", None), ("高", "num"), ("中", "num"), ("低", "num"), ("估值", "num"),
    ("未匹配", "num"), ("无覆盖", "num"), ("单位不匹配", "num"),
    ("有价", "num"), ("命中率", "num"), ("合价(MB)", "num"),
]

# ═══════════════════════════════════════════════════════════════
#  Main entry point
# ═══════════════════════════════════════════════════════════════

def generate_report(config):
    """Generate HTML pricing report. See module docstring for config schema."""
    root = Path(config["root_dir"])
    src_xlsx = root / config["src_xlsx"]
    results_json = root / config["results_json"]
    boq_items_json = root / config["boq_items_json"]
    report_dir = root / "套价报告"
    out_html = report_dir / f"{config['project_name'].replace(' ', '_')}_AI套价报告.html"
    currency = config.get("currency", "THB")
    cur_label = config.get("currency_label", f"百万{currency}")

    # --- Load classification from xlsx ---
    sheet = config.get("sheet_name", "UniqueBQ")
    header_row = config.get("header_row", 5)
    cc = config.get("col_class", {"discipline": 1, "category": 3, "subcategory": 4, "element": 5})
    cq = config.get("col_check", {"desc": 7, "qty": 9})

    wb = openpyxl.load_workbook(str(src_xlsx), data_only=True)
    ws = wb[sheet]
    class_data = {}
    cur = {"discipline": "", "category": "", "subcategory": "", "element": ""}
    for r in range(header_row, ws.max_row + 1):
        if ws.cell(r, cc["discipline"]).value and str(ws.cell(r, cc["discipline"]).value).strip():
            cur["discipline"] = str(ws.cell(r, cc["discipline"]).value).strip()
        if ws.cell(r, cc["category"]).value and str(ws.cell(r, cc["category"]).value).strip():
            cur["category"] = str(ws.cell(r, cc["category"]).value).strip()
        if ws.cell(r, cc["subcategory"]).value and str(ws.cell(r, cc["subcategory"]).value).strip():
            cur["subcategory"] = str(ws.cell(r, cc["subcategory"]).value).strip()
        if ws.cell(r, cc["element"]).value and str(ws.cell(r, cc["element"]).value).strip():
            cur["element"] = str(ws.cell(r, cc["element"]).value).strip()

        desc_val = ws.cell(r, cq["desc"]).value
        qty_val = ws.cell(r, cq["qty"]).value
        if desc_val and isinstance(qty_val, (int, float)) and qty_val > 0:
            class_data[r] = dict(cur)
    wb.close()

    # --- Load results ---
    results = json.loads(results_json.read_text(encoding="utf-8"))
    boq_items = json.loads(boq_items_json.read_text(encoding="utf-8"))
    boq_by_row = {b["excel_row"]: b for b in boq_items}

    for rec in results:
        sup = rec.get("matched_supplier")
        if isinstance(sup, str) and sup.startswith("{"):
            m = re.search(r"'name':\s*'([^']*)'", sup)
            if m:
                rec["matched_supplier"] = m.group(1)

    disc_map = config.get("discipline_map", {})
    disc_order = config.get("discipline_order", [])

    # --- Build rows ---
    rows = []
    for rec in results:
        row = rec["excel_row"]
        boq = boq_by_row.get(row, {})
        cls = class_data.get(row, {})
        disc_en = cls.get("discipline", "")
        disc_cn = disc_map.get(disc_en, disc_en)
        cat = cls.get("category", "")
        subcat = cls.get("subcategory", "")
        elem = cls.get("element", "")

        parts = [p for p in [disc_cn, cat, subcat, elem] if p]
        combined = " > ".join(parts) if parts else (disc_cn or "Unclassified")

        ep = rec.get("converted_price_thb") or rec.get("matched_price_thb")
        try:
            ep = float(ep) if ep is not None else None
        except (ValueError, TypeError):
            ep = None

        mu_raw = rec.get("matched_unit") or ""
        bu_raw = boq.get("unit") or ""
        mu = norm_unit(mu_raw)
        bu = norm_unit(bu_raw)
        orig_price = rec.get("matched_price_thb")
        try:
            op = float(orig_price) if orig_price is not None else None
        except (ValueError, TypeError):
            op = None
        try:
            cp = float(rec.get("converted_price_thb")) if rec.get("converted_price_thb") else None
        except (ValueError, TypeError):
            cp = None
        eff_p = cp or op

        desc_text = boq.get("description_en", "") or ""
        matched_text = rec.get("matched_name") or ""

        conv_note = ""
        if mu != bu and op:
            conv_note = build_conv_note(mu, bu, op, cp, mu_raw, bu_raw, desc_text, matched_text)
        elif cp and op and abs(cp - op) > 0.01:
            conv_note = build_conv_note(mu, bu, op, cp, mu_raw, bu_raw, desc_text, matched_text)

        reasoning_raw = rec.get("reasoning", "") or ""
        reasoning_clean = _clean_reasoning(
            reasoning_raw, boq.get("description_en", ""), rec.get("status", ""),
            matched_name=rec.get("matched_name") or "",
            conv_note=conv_note,
            matched_project=rec.get("matched_project") or "",
            matched_supplier=rec.get("matched_supplier") or "",
        )
        display_unit = bu_raw if (mu != bu and op) else mu_raw

        rows.append({
            "excel_row": row,
            "combined_field": combined,
            "discipline_cn": disc_cn,
            "category": cat,
            "subcategory": subcat,
            "element": elem,
            "description_en": boq.get("description_en", ""),
            "boq_unit": boq.get("unit", ""),
            "boq_qty": boq.get("qty", 0),
            "status": rec.get("status", "no_match"),
            "matched_name": rec.get("matched_name"),
            "matched_unit": mu_raw,
            "display_unit": display_unit,
            "matched_price_thb": rec.get("matched_price_thb"),
            "converted_price_thb": rec.get("converted_price_thb"),
            "matched_project": rec.get("matched_project"),
            "matched_supplier": rec.get("matched_supplier"),
            "matched_date": rec.get("matched_date"),
            "similarity": rec.get("similarity"),
            "eff_price": ep,
            "conv_note": conv_note,
            "reasoning": reasoning_clean,
        })

    rows.sort(key=lambda x: x["excel_row"])

    # --- Stats ---
    h_esc = html.escape
    status_count = Counter(r["status"] for r in rows)
    total = len(rows)
    matched = sum(status_count.get(s, 0) for s in ("high", "medium", "low", "estimated"))
    no_match = status_count.get("no_match", 0)
    known_gap = status_count.get("known_gap", 0)
    unit_gap = status_count.get("unit_gap", 0)
    coverage = round(100 * matched / max(total, 1), 1)

    total_thb = 0.0
    priced = 0
    for r in rows:
        ep_v = r["eff_price"]
        qty = r["boq_qty"] or 0
        if ep_v and qty > 0:
            total_thb += ep_v * qty
            priced += 1

    # --- Discipline breakdown ---
    disc_stats = defaultdict(lambda: {"total": 0, "matched": 0, "total_thb": 0.0,
                                       "high": 0, "medium": 0, "low": 0, "estimated": 0,
                                       "no_match": 0, "known_gap": 0, "unit_gap": 0})
    for r in rows:
        d = r["discipline_cn"] or "Unclassified"
        disc_stats[d]["total"] += 1
        s = r["status"]
        disc_stats[d][s] = disc_stats[d].get(s, 0) + 1
        if s in ("high", "medium", "low", "estimated"):
            disc_stats[d]["matched"] += 1
        ep_v = r["eff_price"]
        qty = r["boq_qty"] or 0
        if ep_v and qty > 0:
            disc_stats[d]["total_thb"] += ep_v * qty

    def _make_summary_row(d_label):
        cnt = disc_stats.get(d_label)
        if not cnt or cnt["total"] == 0:
            return ""
        tot = cnt["total"]
        m = cnt["matched"]
        rate = round(100 * m / max(tot, 1), 1)
        return (
            f'<tr><td><b>{h_esc(d_label)}</b></td>'
            + "".join(f'<td class="num">{cnt.get(k, 0)}</td>'
                      for k in ("high", "medium", "low", "estimated",
                                "no_match", "known_gap", "unit_gap"))
            + f'<td class="num"><b>{m}</b></td>'
            f'<td class="num">{rate}%</td>'
            f'<td class="num">{cnt["total_thb"] / 1e6:,.1f}</td></tr>'
        )

    disc_rows_html = "\n".join(
        row for d in disc_order
        if (row := _make_summary_row(d))
    )

    # --- Detail rows ---
    def _make_th_row(cols):
        return "<tr>" + "".join(
            f'<th{f' class="{cls}"' if cls else ""}>{h_esc(label)}</th>'
            for label, cls in cols
        ) + "</tr>"

    def build_detail_rows(status_list, items, use_gap_header=False):
        out = []
        last_cat = None
        for r in items:
            if r["status"] not in status_list:
                continue
            cat = r.get("category") or ""
            if cat and cat != last_cat:
                last_cat = cat
                disc = r.get("discipline_cn") or ""
                subcat = r.get("subcategory") or ""
                subcat_part = f" > {subcat}" if subcat else ""
                out.append(
                    f'<tr class="cat-header"><td colspan="15">'
                    f'<b>{h_esc(disc)}</b> &gt; <b>{h_esc(cat)}{h_esc(subcat_part)}</b>'
                    f'</td></tr>'
                )

            badge_cls = BADGE_CLASS.get(r["status"], "b-no_match")
            badge_label = STATUS_LABEL.get(r["status"], r["status"])
            badge = f'<span class="badge {badge_cls}">{badge_label}</span>'
            ep_v = r["eff_price"]
            pd = f"{ep_v:,.2f}" if ep_v else "--"

            sup = r.get("matched_supplier") or ""
            if not isinstance(sup, str):
                sup = str(sup)
            if sup.startswith("{") and "name" in sup:
                m = re.search(r"'name':\s*'([^']*)'", sup)
                sup = m.group(1) if m else sup[:30]

            display_unit = r.get("display_unit") or r.get("matched_unit") or ""
            elem = r.get("element") or ""
            mdate = r.get("matched_date") or ""

            out.append(
                f'<tr data-status="{r["status"]}">'
                f'<td class="num">{r["excel_row"]}</td>'
                f'<td class="elem" title="{h_esc(elem)}">{h_esc(elem[:20])}</td>'
                f'<td class="desc" title="{h_esc(r["description_en"])}">{h_esc(r["description_en"][:65])}</td>'
                f'<td class="unit-cell">{h_esc(r["boq_unit"])}</td>'
                f'<td class="num">{r["boq_qty"]:,.1f}</td>'
                f'<td>{badge}</td>'
                f'<td class="desc" title="{h_esc(r["matched_name"] or "")}">{h_esc((r["matched_name"] or "")[:50])}</td>'
                f'<td class="unit-cell">{h_esc(display_unit)}</td>'
                f'<td class="price">{pd}</td>'
                f'<td class="unit-cell">{currency}</td>'
                f'<td class="conv">{h_esc(r["conv_note"])}</td>'
                f'<td class="num">{h_esc(mdate)}</td>'
                f'<td class="desc">{h_esc(sup[:14])}</td>'
                f'<td class="desc">{h_esc((r.get("matched_project") or "")[:14])}</td>'
                f'<td class="reason" title="{h_esc(r["reasoning"])}">{h_esc(r["reasoning"][:240])}</td>'
                f'</tr>'
            )
        return "\n".join(out)

    high_med_rows = build_detail_rows(["high", "medium"], rows)
    low_est_rows = build_detail_rows(["low", "estimated"], rows)
    gap_rows = build_detail_rows(["no_match", "known_gap", "unit_gap"], rows)

    # --- Mermaid diagram ---
    mc = config.get("mermaid", {})
    mermaid_diagram = f"""\
graph LR
    BOQ["&#128203; BOQ Excel&lt;br/&gt;{h_esc(mc.get('boq_label', 'BOQ Items'))}"] --> PY["&#128013; ① Python 路由编排&lt;br/&gt;MongoDB预搜 · 候选检索 · 分包"]
    PY --> HAIKU["&#9889; ② Haiku xN 并行&lt;br/&gt;语义匹配执行 · 独立信号1"]
    HAIKU --> SONNET["&#128270; ③ Sonnet Auditor&lt;br/&gt;独立审核 · 跨模型校验 · 独立信号2"]
    SONNET --> MERGE["&#128202; ④ 合并 + 写回&lt;br/&gt;Excel {h_esc(config.get('writeback_desc', 'Q-Z 列'))} · HTML报告"]
    MONGO["&#128451; {h_esc(mc.get('data_source_label', 'MongoDB 报价库'))}"] -.->|数据源| PY
    RULES["&#128221; {h_esc(mc.get('rules_label', '人工修正规则'))}"] -.->|注入| MERGE
    PY -.->|检索候选| MONGO
    style BOQ fill:#dbeafe,stroke:#2563eb,stroke-width:2px
    style PY fill:#fef3c7,stroke:#d97706,stroke-width:2px
    style HAIKU fill:#d1fae5,stroke:#059669,stroke-width:2px
    style SONNET fill:#ede9fe,stroke:#7c3aed,stroke-width:2px
    style MERGE fill:#fce7f3,stroke:#db2777,stroke-width:2px
    style MONGO fill:#e0e7ff,stroke:#4f46e5,stroke-width:2px
    style RULES fill:#fef2f2,stroke:#dc2626,stroke-width:2px"""

    # --- Assemble HTML ---
    ts_h = datetime.now().strftime("%Y-%m-%d %H:%M")
    out_xlsx_name = config.get("out_xlsx_name", "")
    data_src = h_esc(config.get("data_source_desc", ""))
    writeback = h_esc(config.get("writeback_desc", "Q-Z 列"))
    subtitle_lines = config.get("subtitle_lines", [])
    version_note = h_esc(config.get("version_note", ""))
    rules_text = h_esc(config.get("rules_text", ""))
    human_extra = h_esc(config.get("human_rules_extra", ""))
    proj_name = h_esc(config["project_name"])

    subtitle_html = "\n".join(f"<p>{line}</p>" for line in subtitle_lines)
    alert_html = ""
    if no_match + known_gap + unit_gap > 0:
        alert_html = (
            f'<div class="alert">需人工处理：{no_match} 条未匹配需询价或找专业分包，'
            f'{known_gap} 条已知价格库无覆盖，{unit_gap} 条单位不匹配无法套价。</div>'
        )

    hc = status_count.get("high", 0)
    mc_ = status_count.get("medium", 0)

    html_out = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="UTF-8">
<title>{proj_name} BOQ AI套价报告 -- {ts_h}</title>
{MERMAID_JS}
{MERMAID_INIT}
<style>
{CSS}
</style></head><body><div class="container">

<div class="header">
  <h1>{proj_name} UniqueBQ AI套价报告</h1>
  <p>生成时间：{ts_h} | 数据源：{data_src} | 写回：{out_xlsx_name} {writeback}</p>
  {subtitle_html}
</div>

<div class="cards">
  <div class="card"><div class="num">{total}</div><div class="label">AI 待套价项</div></div>
  <div class="card"><div class="num">{matched}</div><div class="label">有价项</div></div>
  <div class="card"><div class="num">{coverage}%</div><div class="label">命中率</div></div>
  <div class="card"><div class="num">{no_match + known_gap + unit_gap}</div><div class="label">未匹配(需询价)</div></div>
  <div class="card"><div class="num">{total_thb / 1e6:,.1f}</div><div class="label">已套合价({cur_label})</div></div>
  <div class="card"><div class="num">{priced}</div><div class="label">已标价行数</div></div>
</div>

<div class="section"><h2>一、摘要</h2>
  <p>共 <b>{total}</b> 个 AI 待办项：高置信 <b>{hc}</b>、中置信 <b>{mc_}</b>、低置信 <b>{status_count.get("low", 0)}</b>、估值换算 <b>{status_count.get("estimated", 0)}</b>、未匹配 <b>{no_match}</b>、已知无覆盖 <b>{known_gap}</b>、单位不匹配 <b>{unit_gap}</b>。有价项 <b>{matched}</b>，命中率 <b>{coverage}%</b>。</p>
  <p>采用<strong>4层流水线架构</strong>：Python 预搜 MongoDB -> Haiku xN 并行执行（独立信号1）-> Sonnet Auditor 独立审核（独立信号2）-> Merge + Writeback。修正规则：{rules_text}{human_extra}</p>
  <div class="mermaid-wrap"><div class="mermaid">
{mermaid_diagram}
  </div></div>
  {f'<p><strong>本版更新</strong>：{version_note}</p>' if version_note else ''}
  {alert_html}
</div>

<div class="section"><h2>二、按专业拆解</h2>
<table class="summary-table">
{_make_th_row(SUMMARY_HEADER_COLS)}
{disc_rows_html}
</table></div>

<div class="section"><h2>三、高/中置信度匹配明细（{hc + mc_} 条）</h2>
<div class="detail-table"><table>
{_make_th_row(DETAIL_HEADER_COLS)}
{high_med_rows}
</table></div></div>

<div class="section"><h2>四、低置信/估值换算明细（{status_count.get("low", 0) + status_count.get("estimated", 0)} 条）</h2>
<div class="detail-table"><table>
{_make_th_row(DETAIL_HEADER_COLS)}
{low_est_rows}
</table></div></div>

<div class="section"><h2>五、未匹配/已知缺口明细（{no_match + known_gap + unit_gap} 条）</h2>
<div class="note">以下条目需人工处理：no_match = 需重新询价或找专业分包报价；known_gap = 价格库中无此类目；unit_gap = 单位不匹配。</div>
<div class="detail-table"><table>
{_make_th_row(DETAIL_HEADER_COLS_GAP)}
{gap_rows}
</table></div></div>

<div class="footer">
 <p>4层流水线架构 | Python预搜 -> Haiku xN并行 -> Sonnet独立审核 -> Merge回写</p>
 <p>{ts_h}</p>
</div>
</div></body></html>"""

    report_dir.mkdir(parents=True, exist_ok=True)
    out_html.write_text(html_out, encoding="utf-8")
    size_kb = out_html.stat().st_size // 1024
    print(f"HTML report: {out_html}")
    print(f"Size: {size_kb} KB")
    print(f"Total items: {total} | Matched: {matched} ({coverage}%) | Priced: {total_thb / 1e6:,.1f} {cur_label}")

    return {
        "html_path": str(out_html),
        "total": total,
        "matched": matched,
        "coverage": coverage,
        "total_priced": total_thb,
        "priced_rows": priced,
    }


# ═══════════════════════════════════════════════════════════════
#  CLI
# ═══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate AI pricing report HTML")
    parser.add_argument("--config", required=True, help="Path to JSON config file")
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text(encoding="utf-8"))
    generate_report(config)
