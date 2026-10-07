# -*- coding: utf-8 -*-
"""BOQ 清单 -> HTML 审阅稿：编号/英文描述/中文描述/单位 + 审阅意见列；修改与意见保存在内嵌 review-data。"""
import argparse
import datetime
import html
import io
import json
import os
import sys

import openpyxl

FIELDS = ("code", "en", "zh", "unit")
CODE_KEYS = {"no.", "no", "code", "item", "item no.", "编号", "序号"}
UNIT_KEYS = {"unit", "units", "单位"}
FILL_CLASS = {"333F4F": "chapter", "D9E1F2": "subchap", "FBE5D6": "subhead"}
MARK_CLASS = {"C6EFCE": "newrow", "FFFF00": "modrow"}


def read_bytes(path):
    # Excel 打开时以全共享方式持有句柄，普通 open 会 PermissionError
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        k = ctypes.windll.kernel32
        k.CreateFileW.restype = wintypes.HANDLE
        k.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                  wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        h = k.CreateFileW(path, 0x80000000, 7, None, 3, 0, None)
        if h and h != ctypes.c_void_p(-1).value:
            try:
                buf, got, out = ctypes.create_string_buffer(1 << 20), wintypes.DWORD(0), bytearray()
                while k.ReadFile(h, buf, len(buf), ctypes.byref(got), None) and got.value:
                    out += buf.raw[:got.value]
                return bytes(out)
            finally:
                k.CloseHandle(h)
    with open(path, "rb") as f:
        return f.read()


def file_sig(path):
    st = os.stat(path)
    return {"source_mtime": datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
            "source_size": st.st_size}


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def find_header(ws, max_scan=30):
    for row in ws.iter_rows(min_row=1, max_row=min(max_scan, ws.max_row)):
        found = {}
        for c in row:
            t = _text(c.value)
            low = t.lower()
            if not t:
                continue
            if low in CODE_KEYS:
                found.setdefault("code", c.column_letter)
            elif "中文" in t:
                found.setdefault("zh", c.column_letter)
            elif low.startswith("description"):
                found.setdefault("en", c.column_letter)
            elif low in UNIT_KEYS:
                found.setdefault("unit", c.column_letter)
        if {"code", "en", "unit"} <= set(found):
            found.setdefault("zh", None)
            return row[0].row, found
    return None, None


def _rgb(color):
    try:
        if color is not None and color.type == "rgb" and color.rgb:
            return str(color.rgb)[-6:].upper()
    except (AttributeError, TypeError):
        pass
    return None


def read_rows(ws, header_row, cols):
    labels = {f: _text(ws["%s%d" % (cols[f], header_row)].value) for f in FIELDS if cols.get(f)}
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = {f: (_text(ws["%s%d" % (cols[f], r)].value) if cols.get(f) else "") for f in FIELDS}
        if not any(vals.values()) or all(vals[f] == labels[f] for f in labels):
            continue
        en, code = ws["%s%d" % (cols["en"], r)], ws["%s%d" % (cols["code"], r)]
        fill = (_rgb(en.fill.fgColor) if en.fill.patternType else None) or \
               (_rgb(code.fill.fgColor) if code.fill.patternType else None)
        deleted = bool(en.font.strike) or _rgb(en.font.color) == "FF0000"
        rows.append(dict(r=r, fill=fill, deleted=deleted, **vals))
    return rows


def row_class(x):
    if x.get("deleted"):
        return "delrow"
    fill = x.get("fill") or ""
    cls = FILL_CLASS.get(fill, "")
    if not cls:
        t = (x.get("en") or "") + (x.get("zh") or "")
        cls = "chapter" if "【" in t else "subchap" if "《" in t else "subhead" if "{" in t else ""
    return " ".join(c for c in (cls, MARK_CLASS.get(fill, "")) if c)


CSS = r"""<style>
*{box-sizing:border-box}
html,body{height:100%}
body{margin:0;display:flex;flex-direction:column;background:#eef1f5;color:#1f2937;font-family:"Microsoft YaHei","PingFang SC","Segoe UI",sans-serif;font-size:13px;line-height:1.5}
.bar{flex:none;background:#0f172a;color:#fff;padding:9px 14px;display:flex;gap:8px;align-items:center;flex-wrap:wrap}
.bar b{font-size:15px;margin-right:6px}
.bar button{background:#2563eb;color:#fff;border:0;border-radius:6px;padding:6px 11px;cursor:pointer;font-size:12.5px}
.bar button:hover{background:#1d4ed8}
.bar button.warn{background:#b91c1c}
.badge{background:#f59e0b;color:#111;border-radius:11px;padding:1px 9px;font-size:12px;font-weight:700}
.legend{flex:none;display:flex;gap:14px;flex-wrap:wrap;margin:9px 14px;font-size:12.5px;color:#334155}
.legend i{display:inline-block;width:14px;height:14px;border-radius:3px;vertical-align:-2px;margin-right:5px;border:1px solid #cbd5e1}
.wrap{flex:1 1 auto;min-height:0;overflow:auto;background:#fff}
table{border-collapse:separate;border-spacing:0;table-layout:fixed}
th,td{border-right:1px solid #dfe3e8;border-bottom:1px solid #e8ecf1;padding:5px 8px;vertical-align:top;text-align:left;overflow-wrap:anywhere}
thead th{position:sticky;top:0;z-index:10;background:#334155;color:#fff;text-align:center;font-weight:700}
th .rz{position:absolute;top:0;right:-3px;width:7px;height:100%;cursor:col-resize;z-index:11}
th .rz:hover{background:#93c5fd}
td.no{white-space:nowrap;font-family:Consolas,monospace;font-size:12px;color:#1d4ed8}
td.tx,td.note{white-space:pre-wrap}
td.zh{color:#334155}
td.unit{text-align:center;white-space:nowrap}
td.rn{text-align:right;color:#94a3b8;font-family:Consolas,monospace;font-size:11px;background:#f8fafc}
tr.chapter td{background:#333F4F;color:#fff;font-weight:700}
tr.chapter td.no{color:#c7d2fe}
tr.subchap td{background:#D9E1F2;font-weight:600}
tr.subhead td{background:#FBE5D6;font-weight:600}
tr.newrow td{background:#C6EFCE}
tr.modrow td{background:#FFFF00}
tr.delrow td{background:#FFE1E1;color:#B91C1C;text-decoration:line-through}
tr td.note{background:#fffbeb;color:#92400e;text-decoration:none;font-weight:400}
td.rn{cursor:pointer}
td del{color:#dc2626;text-decoration:line-through}
td ins{color:#2563eb;text-decoration:underline}
tr.rowdel td:not(.note){background:#fee2e2 !important;color:#dc2626 !important;text-decoration:line-through}
tr.noted td.rn{box-shadow:inset 3px 0 0 #d97706;color:#b45309;font-weight:700}
td[contenteditable]:focus{background:#fff;outline:2px solid #2563eb}
td.edited{box-shadow:inset 3px 0 0 #ea580c,inset 0 0 0 1px #fdba74}
.tip{position:fixed;right:14px;bottom:14px;background:#0f172a;color:#fff;padding:8px 12px;border-radius:8px;font-size:12.5px;opacity:0;transition:opacity .2s;z-index:30}
.tip.on{opacity:.95}
</style>"""

JS = r"""<script>
var DATA=JSON.parse(document.getElementById('review-data').textContent);
var LS_KEY='pkreview_'+DATA.meta.source+'|'+DATA.meta.source_mtime;
var edits={}, notes={}, dels={};
function norm(s){return (s||'').replace(/\r\n?/g,'\n').replace(/\u00a0/g,' ').replace(/[ \t]+\n/g,'\n').trim();}
function esc(s){return s.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');}
function tip(t){var e=document.getElementById('tip');e.textContent=t;e.classList.add('on');clearTimeout(e._t);e._t=setTimeout(function(){e.classList.remove('on');},3000);}
function vals(o){return Object.keys(o).map(function(k){return o[k];}).sort(function(a,b){return (a.row||1e9)-(b.row||1e9);});}
function updCnt(){['E','N','D'].forEach(function(k,i){document.getElementById('cnt'+k).textContent=Object.keys([edits,notes,dels][i]).length;});}
function saveLS(){try{localStorage.setItem(LS_KEY,JSON.stringify({edits:vals(edits),comments:vals(notes),deletes:vals(dels)}));}catch(e){}}
function info(tr){var r=tr.dataset.r,z=tr.querySelector('td[data-c="zh"]');return {rid:r,row:/^\d+$/.test(r)?parseInt(r,10):null,code:tr.dataset.code,desc:z?norm(z.dataset.orig):''};}
// 字符级修订显示：删掉的字红色删除线，新加的字蓝色下划线
function diffHTML(a,b){
  var n=a.length,m=b.length;
  if(n*m>4000000)return '<del>'+esc(a)+'</del><ins>'+esc(b)+'</ins>';
  var D=[];for(var i=0;i<=n;i++){D.push(new Uint16Array(m+1));}
  for(i=n-1;i>=0;i--)for(var j=m-1;j>=0;j--)D[i][j]=a[i]===b[j]?D[i+1][j+1]+1:Math.max(D[i+1][j],D[i][j+1]);
  var out='',i2=0,j2=0,mode='',buf='';
  function emit(k,ch){if(k!==mode){out+=wrap(mode,buf);mode=k;buf='';}buf+=ch;}
  function wrap(k,s){return !s?'':k==='d'?'<del>'+esc(s)+'</del>':k==='i'?'<ins>'+esc(s)+'</ins>':esc(s);}
  while(i2<n&&j2<m){if(a[i2]===b[j2]){emit('e',a[i2]);i2++;j2++;}else if(D[i2+1][j2]>=D[i2][j2+1]){emit('d',a[i2]);i2++;}else{emit('i',b[j2]);j2++;}}
  while(i2<n){emit('d',a[i2]);i2++;}
  while(j2<m){emit('i',b[j2]);j2++;}
  return out+wrap(mode,buf);
}
function show(td){var o=norm(td.dataset.orig),c=td.dataset.cur;if(c===undefined||c===o){td.textContent=td.dataset.orig;}else{td.innerHTML=diffHTML(o,c);}}
function track(td){
  var tr=td.closest('tr'), c=td.dataset.c, cur=td.dataset.cur===undefined?norm(td.dataset.orig):td.dataset.cur, x=info(tr);
  if(c==='note'){
    if(cur){x.text=cur;notes[x.rid]=x;tr.classList.add('noted');}else{delete notes[x.rid];tr.classList.remove('noted');}
  }else{
    var o=norm(td.dataset.orig),k=x.rid+'|'+c;
    if(cur!==o){x.col=c;x.old=o;x.new=cur;delete x.desc;edits[k]=x;td.classList.add('edited');}
    else{delete edits[k];td.classList.remove('edited');delete td.dataset.cur;}
  }
}
document.querySelectorAll('td[contenteditable]').forEach(function(td){
  td.addEventListener('focus',function(){if(td.dataset.c!=='note'&&td.dataset.cur!==undefined)td.textContent=td.dataset.cur;});
  td.addEventListener('input',function(){td.dataset.cur=norm(td.innerText);track(td);updCnt();saveLS();});
  td.addEventListener('blur',function(){if(td.dataset.c!=='note')show(td);});
});
function setDel(tr,on){var x=info(tr);if(on){dels[x.rid]=x;tr.classList.add('rowdel');}else{delete dels[x.rid];tr.classList.remove('rowdel');}}
document.querySelectorAll('#t tbody td.rn').forEach(function(td){td.title='点击标记/取消整行删除';td.addEventListener('click',function(){var tr=td.closest('tr');setDel(tr,!tr.classList.contains('rowdel'));updCnt();saveLS();});});
function put(rid,c,text){var td=document.querySelector('tr[data-r="'+rid+'"] td[data-c="'+c+'"]');if(!td)return;td.dataset.cur=text;track(td);if(c==='note')td.textContent=text;else show(td);}
function restore(){
  var st=DATA.review;
  try{var s=localStorage.getItem(LS_KEY);if(s)st=JSON.parse(s);}catch(e){}
  (st.edits||[]).forEach(function(x){put(x.rid,x.col,x.new);});
  (st.comments||[]).forEach(function(x){put(x.rid,'note',x.text);});
  (st.deletes||[]).forEach(function(x){var tr=document.querySelector('tr[data-r="'+x.rid+'"]');if(tr)setDel(tr,true);});
  updCnt();
}
function clearAll(){
  if(!confirm('清除全部修改、意见和删除标记，恢复到生成时的内容？'))return;
  document.querySelectorAll('td[contenteditable]').forEach(function(td){delete td.dataset.cur;td.textContent=td.dataset.orig;td.classList.remove('edited');});
  document.querySelectorAll('tr.noted,tr.rowdel').forEach(function(tr){tr.classList.remove('noted','rowdel');});
  edits={};notes={};dels={};saveLS();updCnt();tip('已清除；保存后审阅稿里的记录也会清空');
}
function copyChanges(){
  var T='\t',L=[];
  vals(dels).forEach(function(x){L.push(['删除行',x.row||'新',x.code,'',x.desc,''].join(T));});
  vals(edits).forEach(function(x){L.push(['修改',x.row||'新',x.code,x.col,x.old,x.new].join(T).replace(/\n/g,' '));});
  vals(notes).forEach(function(x){L.push(['意见',x.row||'新',x.code,'',x.text,''].join(T).replace(/\n/g,' '));});
  if(!L.length){tip('没有修改、意见或删除');return;}
  navigator.clipboard.writeText(['类型','行','编号','列','原值/意见','新值'].join(T)+'\n'+L.join('\n')).then(function(){tip('已复制 '+L.length+' 条');});
}
function buildHTML(){
  var doc=document.documentElement.cloneNode(true);
  doc.querySelectorAll('td[contenteditable]').forEach(function(td){td.textContent=td.dataset.orig;td.classList.remove('edited');td.removeAttribute('data-cur');});
  doc.querySelectorAll('tr.noted,tr.rowdel').forEach(function(tr){tr.classList.remove('noted','rowdel');});
  doc.querySelector('#tip').className='tip';
  doc.querySelector('#review-data').textContent=JSON.stringify(DATA).replace(/<\//g,'<\\/');
  return '<!doctype html>\n'+doc.outerHTML;
}
// 记住第一次选定的文件句柄，之后保存只询问覆盖
var HKEY=DATA.meta.source+'|'+DATA.meta.out_name;
function idb(){return new Promise(function(res,rej){var r=indexedDB.open('pkreview',1);r.onupgradeneeded=function(){r.result.createObjectStore('h');};r.onsuccess=function(){res(r.result);};r.onerror=function(){rej(r.error);};});}
async function getH(){try{var db=await idb();return await new Promise(function(res){var q=db.transaction('h').objectStore('h').get(HKEY);q.onsuccess=function(){res(q.result||null);};q.onerror=function(){res(null);};});}catch(e){return null;}}
async function setH(h){try{var db=await idb();db.transaction('h','readwrite').objectStore('h').put(h,HKEY);}catch(e){}}
async function saveFile(){
  var rv={saved_at:new Date().toISOString(),edits:vals(edits),comments:vals(notes),deletes:vals(dels)};
  var n='（修改 '+rv.edits.length+' · 意见 '+rv.comments.length+' · 删除 '+rv.deletes.length+'）', name=DATA.meta.out_name;
  if(!window.showSaveFilePicker){
    DATA.review=rv;var a=document.createElement('a');a.href=URL.createObjectURL(new Blob([buildHTML()],{type:'text/html'}));a.download=name;a.click();
    tip('浏览器不支持直接保存，已下载到「下载」文件夹：'+name+n);return;
  }
  var h=await getH();
  if(h&&!confirm('覆盖保存到 '+h.name+'？\n点"取消"可另选位置。'))h=null;
  if(h){try{if((await h.queryPermission({mode:'readwrite'}))!=='granted'&&(await h.requestPermission({mode:'readwrite'}))!=='granted')h=null;}catch(e){h=null;}}
  if(!h){
    try{h=await window.showSaveFilePicker({suggestedName:name,types:[{description:'HTML 审阅稿',accept:{'text/html':['.html']}}]});}
    catch(e){tip('已取消保存');return;}
    await setH(h);
  }
  DATA.review=rv;
  try{var w=await h.createWritable();await w.write(buildHTML());await w.close();tip('已保存：'+h.name+n);}
  catch(e){tip('保存失败：'+(e&&e.message||e));}
}
var cols=document.querySelectorAll('#t colgroup col'), W=[];
function applyW(){cols.forEach(function(c,i){c.style.width=W[i]+'px';});document.getElementById('t').style.width=W.reduce(function(a,b){return a+b;},0)+'px';}
function computeInit(){var rest=document.querySelector('.wrap').clientWidth-2-56-120-70;if(rest<400)rest=1000;W=[56,120,Math.round(rest*0.40),Math.round(rest*0.34),70,Math.round(rest*0.26)];applyW();}
function resetWidths(){try{localStorage.removeItem(LS_KEY+'|w');}catch(e){}computeInit();tip('列宽已重置');}
document.querySelectorAll('#t thead th .rz').forEach(function(rz,i){
  rz.addEventListener('mousedown',function(e){
    e.preventDefault();var x0=e.pageX,w0=W[i],sum=W[i]+W[i+1];document.body.style.userSelect='none';
    function mv(ev){var a=Math.max(40,Math.min(w0+ev.pageX-x0,sum-40));W[i]=a;W[i+1]=sum-a;applyW();}
    function up(){document.removeEventListener('mousemove',mv);document.removeEventListener('mouseup',up);document.body.style.userSelect='';try{localStorage.setItem(LS_KEY+'|w',JSON.stringify(W));}catch(e){}}
    document.addEventListener('mousemove',mv);document.addEventListener('mouseup',up);
  });
});
try{var sw=JSON.parse(localStorage.getItem(LS_KEY+'|w'));if(sw&&sw.length===6){W=sw;applyW();}else computeInit();}catch(e){computeInit();}
restore();
</script>"""

LEGEND = ('<div class="legend">'
          '<span><i style="background:#333F4F"></i>章【】</span><span><i style="background:#D9E1F2"></i>节《》</span>'
          '<span><i style="background:#FBE5D6"></i>小标题{}</span><span><i style="background:#FFFF00"></i>调整</span>'
          '<span><i style="background:#C6EFCE"></i>新增</span><span><i style="background:#FFE1E1"></i>删除</span>'
          '<span><i style="background:#fdba74"></i>本次修改</span><span><i style="background:#fffbeb"></i>审阅意见</span>'
          '<span style="color:#64748b">单元格可直接改（删字红色删除线、加字蓝色下划线）· 点行号标记整行删除 · 最右列写意见 · 改完点"保存"</span></div>')


def _e(v, quote=False):
    return html.escape(v or "", quote=quote)


def render(rows, dst, title, meta):
    data = {"meta": meta, "review": {"saved_at": None, "edits": [], "comments": []}}
    H = ['<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">',
         '<title>%s</title>' % _e(title), CSS, '</head><body>',
         '<div class="bar"><b>%s</b>' % _e(title),
         '<button onclick="saveFile()">保存</button><button onclick="copyChanges()">复制变更清单</button>',
         '<button onclick="resetWidths()">重置列宽</button><button class="warn" onclick="clearAll()">清除全部编辑</button>',
         '<span>修改 <span class="badge" id="cntE">0</span> 处 · 意见 <span class="badge" id="cntN">0</span> 条 · 删除 <span class="badge" id="cntD">0</span> 行</span></div>',
         LEGEND,
         '<div class="wrap"><table id="t"><colgroup>' + '<col>' * 6 + '</colgroup><thead><tr>'
         '<th>行<span class="rz"></span></th><th>编号<span class="rz"></span></th><th>Description<span class="rz"></span></th>'
         '<th>中文描述<span class="rz"></span></th><th>单位<span class="rz"></span></th><th>审阅意见</th></tr></thead><tbody>']
    for n, x in enumerate(rows, 1):
        r = x.get("r") or 0
        fill, cls = x.get("fill"), row_class(x)
        style = ' style="background:#%s"' % fill if fill and fill not in FILL_CLASS and fill not in MARK_CLASS and cls != "delrow" else ""
        H.append('<tr class="%s" data-r="%s" data-code="%s"%s><td class="rn">%s</td>'
                 % (cls, r or "n%d" % n, _e(x.get("code"), True), style, r or "新"))
        for f, css in (("code", "no"), ("en", "tx"), ("zh", "tx zh"), ("unit", "unit")):
            H.append('<td class="%s" contenteditable="true" data-c="%s" data-orig="%s">%s</td>'
                     % (css, f, _e(x.get(f), True), _e(x.get(f))))
        H.append('<td class="note" contenteditable="true" data-c="note" data-orig=""></td></tr>')
    H.append('</tbody></table></div><div class="tip" id="tip"></div>')
    H.append('<script type="application/json" id="review-data">%s</script>'
             % json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
    H += [JS, '</body></html>']
    with io.open(dst, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(H))
    return len(rows)


def build(src, sheet=None, header_row=None):
    src = os.path.abspath(src)
    wb = openpyxl.load_workbook(io.BytesIO(read_bytes(src)), data_only=True)
    for name in ([sheet] if sheet else wb.sheetnames):
        ws = wb[name]
        hr, cols = find_header(ws)
        if cols:
            break
    else:
        raise SystemExit("未识别到表头（需要同一行里有 No./编号、Description、Unit/单位）。可用 --sheet 指定工作表。")
    hr = header_row or hr
    stem = os.path.splitext(os.path.basename(src))[0]
    dst = os.path.join(os.path.dirname(src), stem + "_审阅稿.html")
    meta = dict(source=src, sheet=ws.title, header_row=hr, cols=cols, out_name=os.path.basename(dst), **file_sig(src))
    n = render(read_rows(ws, hr, cols), dst, "%s · %s · 审阅稿" % (stem, ws.title), meta)
    return dst, n


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="BOQ xlsx -> HTML 审阅稿")
    ap.add_argument("xlsx")
    ap.add_argument("--sheet")
    ap.add_argument("--header-row", type=int)
    a = ap.parse_args(argv)
    dst, n = build(a.xlsx, a.sheet, a.header_row)
    print("审阅稿：%s（%d 行）" % (dst, n))


if __name__ == "__main__":
    main()
