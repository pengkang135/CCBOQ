import json
import os
import re
import sys

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))
import review_html  # noqa: E402
import review_read  # noqa: E402

ROWS = [
    ("A.01", "【Earthworks】", "【土方】", "", "333F4F", None),
    ("A.01.01", "《Excavation》", "《开挖》", "", "D9E1F2", None),
    ("A.01.01.001", "Trench excavation\ndepth ≤2m", "沟槽开挖\n深度≤2m", "m3", None, None),
    ("A.01.01.002", 'Backfill a<b & "c" </script>', "回填", "m3", "FFFF00", None),
    ("A.01.01.003", "Old item", "旧项目", "m2", None, "del"),
    (None,) * 6,
    ("A.01.01.004", "Formula row", "公式行", "nr", None, "formula"),
]


def make_boq(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BQ"
    ws["B1"] = "Fixture BOQ"
    header = ("No.", "Description", "中文描述", "Unit")
    for i, v in enumerate(header):
        ws.cell(3, 2 + i, v)
    for r, (code, en, zh, unit, rgb, flag) in enumerate(ROWS, 4):
        if code is None:
            continue
        for i, v in enumerate((code, en, zh, unit)):
            ws.cell(r, 2 + i, v)
        if rgb:
            ws.cell(r, 3).fill = PatternFill("solid", fgColor=rgb)
        if flag == "del":
            ws.cell(r, 3).font = Font(color="FF0000", strike=True)
        if flag == "formula":
            ws.cell(r, 5).value = "=E6"
    for i, v in enumerate(header):
        ws.cell(4 + len(ROWS), 2 + i, v)
    wb.save(path)


@pytest.fixture
def built(tmp_path):
    src = str(tmp_path / "Fixture BOQ_v1.xlsx")
    make_boq(src)
    dst, n = review_html.build(src)
    return src, dst, n, open(dst, encoding="utf-8").read()


def test_header_rows_and_meta(built):
    src, dst, n, t = built
    assert os.path.basename(dst) == "Fixture BOQ_v1_审阅稿.html"
    assert n == 6
    assert re.findall(r'<tr class="([^"]*)" data-r="(\d+)"', t) == [
        ("chapter", "4"), ("subchap", "5"), ("", "6"), ("modrow", "7"), ("delrow", "8"), ("", "10")]
    meta = review_read.load(dst)["meta"]
    assert (meta["header_row"], meta["cols"]) == (3, {"code": "B", "en": "C", "zh": "D", "unit": "E"})
    assert t.count("<button") == 4 and t.count('data-c="note"') == 6


def test_escaping_newline_formula(built):
    _, _, _, t = built
    assert 'data-orig="Backfill a&lt;b &amp; &quot;c&quot; &lt;/script&gt;"' in t
    assert 'data-orig="沟槽开挖\n深度≤2m"' in t
    assert "=E6" not in t


def test_read_back(built, capsys):
    src, dst, _, t = built
    m = review_read.DATA_RE.search(t)
    d = json.loads(m.group(1))
    d["review"] = {"saved_at": "2026-10-06T10:00:00Z",
                   "edits": [{"rid": "7", "row": 7, "code": "A.01.01.002", "col": "en",
                              "old": 'Backfill a<b & "c" </script>', "new": "Backfill </script> fixed"}],
                   "comments": [{"rid": "8", "row": 8, "code": "A.01.01.003", "desc": "旧项目", "text": "整行删除"}],
                   "deletes": [{"rid": "10", "row": 10, "code": "A.01.01.004", "desc": "公式行"}]}
    with open(dst, "w", encoding="utf-8") as f:
        f.write(t[:m.start(1)] + json.dumps(d, ensure_ascii=False).replace("</", "<\\/") + t[m.end(1):])
    review_read.main([dst])
    out = capsys.readouterr().out
    assert "Backfill </script> fixed" in out and "整行删除" in out
    assert "| 10 | A.01.01.004 | 公式行 |" in out
    assert "合计：删除 1 行，修改 1 处，意见 1 条" in out and "注意" not in out
    os.utime(src, (0, 0))
    review_read.main([dst])
    assert "审阅期间源文件被改过" in capsys.readouterr().out
