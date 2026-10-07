"""清单价格回填：按匹配 JSON 把源报价写入目标表预留列，总额对齐校验。

人工裁决走 --decisions decisions.json，不写在本脚本里——裁决绑定具体项目的编号体系，
硬编码进共享脚本会在换项目时静默错配（源A→目标B 的偏移配对只命中一次不触发重复行检测）。
不传 --decisions 即「无人工裁决」，正是新项目第一次跑的正确初值。
"""
import json
import shutil
import logging
import argparse
from pathlib import Path

# BOQ 混合类型列会触发 fastexcel 的 dtype 回退警告，逐格读取本就按字符串处理，噪音无意义
logging.getLogger("fastexcel").setLevel(logging.ERROR)

import fastexcel
import openpyxl
from openpyxl.styles import PatternFill

GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

DECISION_KEYS = ('manual_fill', 'green_fill', 'green_fill_offset', 'green_insert',
                 'insert_desc', 'item0_desc', 'skip_src', 'special')

SRC_MONEY_COLS = ('rate', 'labour', 'plant', 'mat', 'sub', 'others',
                  'offsite', 'headoffice', 'total')


def load_decisions(path, target):
    """读人工裁决表。未提供则全空——新项目首跑即此状态"""
    empty = {'manual_fill': [], 'green_fill': [], 'green_fill_offset': [],
             'green_insert': [], 'insert_desc': {}, 'item0_desc': {},
             'skip_src': set(), 'special': []}
    if not path:
        return empty

    raw = json.load(open(path, encoding='utf-8'))
    unknown = [k for k in raw if not k.startswith('_') and k not in DECISION_KEYS]
    if unknown:
        raise SystemExit(f'[错误] {path} 含未知键 {unknown}，可用键: {list(DECISION_KEYS)}')

    # 裁决表绑定具体项目的编号体系，张冠李戴会配出毫不相干的两项
    bound = raw.get('_target')
    if bound and Path(bound).name != Path(target).name:
        print(f'!! 警告: {Path(path).name} 声明绑定 {bound!r}，'
              f'但本次目标是 {Path(target).name!r}')
        print('   裁决表跨项目复用会把不相干的条目强行配对，务必核对 dry-run 差额')

    d = dict(empty)
    for k in DECISION_KEYS:
        if k not in raw:
            continue
        v = raw[k]
        if k == 'skip_src':
            d[k] = set(v)
        elif k in ('insert_desc', 'item0_desc'):
            d[k] = dict(v)
        elif k in ('manual_fill', 'green_fill_offset', 'special'):
            d[k] = [tuple(x) for x in v]
        else:
            d[k] = list(v)
    return d


def is_num(v):
    if v in ('nan', '', 'None', None):
        return False
    try:
        float(str(v).replace(',', ''))
        return True
    except (ValueError, AttributeError, TypeError):
        return False


def to_num(v):
    v = str(v).strip()
    if v in ('nan', '', 'None'):
        return None
    try:
        return float(v.replace(',', ''))
    except ValueError:
        return None


def parse_cols(spec, required=()):
    """解析 1-based 列映射，返回 0-based 索引"""
    try:
        cols = {k.strip(): int(v) for k, v in (kv.split('=') for kv in spec.split(','))}
    except ValueError:
        raise SystemExit(f'[错误] 列映射格式应为 key=列号,key=列号：{spec}')
    bad = [k for k, v in cols.items() if v < 1]
    if bad:
        raise SystemExit(f'[错误] 列号是 1-based，{bad} 收到了 <1 的值。A 列=1')
    missing = [k for k in required if k not in cols]
    if missing:
        raise SystemExit(f'[错误] 列映射缺少必需键 {missing}：{spec}')
    return {k: v - 1 for k, v in cols.items()}


def prefix_get(m, key):
    for k, v in m.items():
        if k.startswith(key):
            return v
    return None


def read_source(path, sheet, cols):
    reader = fastexcel.read_excel(path)
    name = pick_sheet(reader, sheet, path)
    df = reader.load_sheet(name, header_row=None).to_pandas()

    ncol = df.shape[1]
    over = {k: v + 1 for k, v in cols.items() if v >= ncol}
    if over:
        raise SystemExit(f'[错误] 源表 {name} 只有 {ncol} 列，但 {over} 指向更右。'
                         f'列号是 1-based，检查是否误传了 0-based 值')

    item_map, desc_map, item0, unit_empty = {}, {}, {}, {}
    for ri in range(len(df)):
        it = str(df.iloc[ri, cols['item']]).strip()
        unit = str(df.iloc[ri, cols['unit']]).strip()
        qty = str(df.iloc[ri, cols['qty']]).strip()
        d = {'desc': str(df.iloc[ri, cols['desc']]).strip(), 'unit': unit,
             'qty': to_num(df.iloc[ri, cols['qty']])}
        for label in SRC_MONEY_COLS:
            d[label] = to_num(df.iloc[ri, cols[label]]) if label in cols else None

        has_unit = unit not in ('nan', '', 'None')
        if it in ('nan', '', 'None'):
            if has_unit and is_num(qty):
                desc_map[d['desc']] = d
        elif it == '0':
            item0[d['desc']] = d
        elif has_unit and is_num(qty):
            item_map[it] = d
        elif is_num(qty):
            unit_empty[it] = d

    if not item_map and not desc_map:
        raise SystemExit(f'[错误] 源表 {name} 未读到任何叶子项。'
                         f'--src-cols 列号（1-based）可能指错列')
    return item_map, desc_map, item0, unit_empty


def pick_sheet(reader, sheet, path):
    if sheet is None:
        return reader.sheet_names[0]
    if sheet not in reader.sheet_names:
        raise SystemExit(f'[错误] {path} 无 sheet {sheet!r}。'
                         f'可用: {", ".join(reader.sheet_names)}')
    return sheet


def read_target(path, key_cols, probe_col):
    """probe_col: 预留列区的 qty 列（1-based），用于放宽叶子判定"""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ci = {k: v + 1 for k, v in key_cols.items()}   # openpyxl 的 cell() 是 1-based

    item_row, empty_item_desc_row, leaf_item_desc = {}, {}, {}
    for r in range(1, ws.max_row + 1):
        it = str(ws.cell(r, ci['item']).value).strip() if ws.cell(r, ci['item']).value else ''
        desc = str(ws.cell(r, ci['desc']).value).strip() if ws.cell(r, ci['desc']).value else ''
        if it and it not in ('nan', 'Item', 'None'):
            unit = ws.cell(r, ci['unit']).value
            trueqty = ws.cell(r, ci['trueqty']).value
            # 陷阱② 叶子判定漏判：插入行只填部分列，trueqty 空，放宽到「trueqty 或预留 qty 列有数」
            is_leaf = (unit not in (None, '') and str(unit).strip() not in ('nan', 'None')
                       and (is_num(trueqty) or is_num(ws.cell(r, probe_col).value)))
            if is_leaf:
                item_row[it] = r
                leaf_item_desc[(it, desc)] = r
            else:
                item_row.setdefault(it, r)
        elif desc:
            empty_item_desc_row.setdefault(desc, r)

    if not item_row:
        raise SystemExit(f'[错误] 目标表未读到任何条目。'
                         f'--tgt-key-cols 列号（1-based）可能指错列')
    return wb, ws, item_row, empty_item_desc_row, leaf_item_desc


def write_target(ws, row, d, tgt):
    for label in SRC_MONEY_COLS:
        if label in tgt:
            ws.cell(row, tgt[label] + 1).value = d.get(label)
    if 'qty' in tgt:
        ws.cell(row, tgt['qty'] + 1).value = d['qty']


def main():
    ap = argparse.ArgumentParser(
        description='按匹配 JSON 回填源报价到目标表预留列',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
所有列号均为 1-based（A 列 = 1），与 compare_boq.py / mark_boq_three_color.py 一致。

decisions.json 结构（全部可选，缺省即空）:
  manual_fill        [[源编号, 目标编号], ...]   显式配对，覆盖自动匹配
  green_fill         [源编号, ...]               源叶子填到目标同编号的标题行，标绿
  green_fill_offset  [[源编号, 目标编号], ...]   同上但编号有偏移
  green_insert       [源编号, ...]               源有目标无，填到目标已存在的对应行
  insert_desc        {源编号: 目标desc前缀}      编号重复时用 desc 消歧定位
  item0_desc         {源编号: 源desc前缀}        源 item 误填为 '0' 时的兜底
  skip_src           [源编号, ...]               跳过自动匹配的误配项
  special            [[src来源, src键, tgt类型, tgt键], ...]
                     src来源: item / desc / unit_empty ; tgt类型: item / desc

示例:
  python fill_boq_prices.py --source 报价.xlsx --target 测算表.xlsx --match match.json \\
      --src-cols item=2,desc=3,unit=4,qty=5,rate=7,labour=8,plant=9,mat=10,sub=11,\\
others=12,offsite=13,headoffice=14,total=15 \\
      --tgt-cols qty=17,rate=18,labour=19,plant=20,mat=21,sub=22,others=23,\\
offsite=24,headoffice=25,total=26 \\
      --decisions decisions.json --expect-total 259150.0 --dry-run
        """)
    ap.add_argument('--source', required=True, help='源报价清单 xlsx')
    ap.add_argument('--target', required=True, help='目标清单 xlsx（就地回填，自动备份）')
    ap.add_argument('--match', required=True, help='match_boq_prices.py 产出的 JSON')
    ap.add_argument('--decisions', default=None,
                    help='人工裁决 JSON，跟项目走。不传即无裁决（新项目首跑）')
    ap.add_argument('--sheet', default=None, help='源表 sheet 名（默认第一个）')
    ap.add_argument('--src-cols', required=True,
                    help='源列映射 1-based，需含 item,desc,unit,qty 及各成本科目')
    ap.add_argument('--tgt-cols', required=True,
                    help='目标预留列区 1-based，如 qty=17,rate=18,...,total=26')
    ap.add_argument('--tgt-key-cols', default='item=2,desc=3,unit=4,trueqty=5',
                    help='目标表键列 1-based（默认 item=2,desc=3,unit=4,trueqty=5）')
    ap.add_argument('--expect-total', type=float, default=None,
                    help='期望总额。写盘必须传且差额为 0；--dry-run 时可省')
    ap.add_argument('--dry-run', action='store_true', help='只统计不写文件')
    args = ap.parse_args()

    src_cols = parse_cols(args.src_cols, required=('item', 'desc', 'unit', 'qty'))
    tgt_cols = parse_cols(args.tgt_cols, required=('qty', 'total'))
    key_cols = parse_cols(args.tgt_key_cols, required=('item', 'desc', 'unit', 'trueqty'))
    dec = load_decisions(args.decisions, args.target)

    if not args.dry_run and args.expect_total is None:
        raise SystemExit('[错误] 写盘必须传 --expect-total 做总额对齐；'
                         '只想看统计请加 --dry-run')

    item_map, desc_map, item0, unit_empty = read_source(args.source, args.sheet, src_cols)
    wb, ws, item_row, empty_item_desc_row, leaf_item_desc = read_target(
        args.target, key_cols, tgt_cols['qty'] + 1)

    fill = [(x['src'], x['tgt']) for x in json.load(open(args.match, encoding='utf-8'))['fill']]
    fill += list(dec['manual_fill'])

    fill_ok = fill_skip = fill_item0 = 0
    skipped_list = []
    planned = []

    def push(row, d, label):
        planned.append((row, d, label))

    for si, ti in fill:
        if si in dec['skip_src']:
            skipped_list.append(f'{si}(skip_src)')
            continue
        if si in item_map:
            r = item_row.get(ti)
            if r is None:
                skipped_list.append(f'{si}(目标无 {ti})')
                continue
            push(r, item_map[si], f'fill {si}->{ti}')
            fill_ok += 1
        elif si in dec['item0_desc']:
            d = prefix_get(item0, dec['item0_desc'][si])
            r = item_row.get(ti)
            if d is None:
                skipped_list.append(f'{si}(item0 desc 未命中)')
            elif r is None:
                skipped_list.append(f'{si}(目标无 {ti})')
            else:
                push(r, d, f'item0 {si}->{ti}')
                fill_item0 += 1
        else:
            skipped_list.append(si)
            fill_skip += 1

    for si, ti in list(dec['green_fill_offset']) + [(x, x) for x in dec['green_fill']]:
        if si in item_map:
            r = item_row.get(ti)
            if r is None:
                skipped_list.append(f'{si}(目标无 {ti})')
                continue
            push(r, item_map[si], f'green {si}->{ti}')
            fill_ok += 1
        else:
            skipped_list.append(si)
            fill_skip += 1

    for si in dec['green_insert']:
        if si in item_map:
            # 陷阱① 重复编号：插入行与原行共享编号，item_row 无法区分，用 desc 消歧
            r = (leaf_item_desc.get((si, dec['insert_desc'][si]))
                 if si in dec['insert_desc'] else item_row.get(si))
            if r is None:
                skipped_list.append(f'{si}(目标无)')
                continue
            push(r, item_map[si], f'insert {si}')
            fill_ok += 1
        else:
            skipped_list.append(si)
            fill_skip += 1

    src_pool = {'item': item_map, 'desc': desc_map, 'unit_empty': unit_empty}
    for src_source, src_key, tgt_type, tgt_key in dec['special']:
        if src_source not in src_pool:
            raise SystemExit(f'[错误] special 的 src来源 {src_source!r} 无效，'
                             f'可用: item / desc / unit_empty')
        d = prefix_get(src_pool[src_source], src_key)
        if d is None:
            skipped_list.append(f'special[{src_key}] 源未命中')
            continue
        r = item_row.get(tgt_key) if tgt_type == 'item' else empty_item_desc_row.get(tgt_key)
        if r is None:
            skipped_list.append(f'special[{src_key}] 目标未命中')
            continue
        push(r, d, f'special {src_key}->{tgt_key}')
        fill_ok += 1

    # 陷阱⑤ dry-run 与写回不一致：重复行 double-count，写盘前必须拦住
    from collections import Counter
    row_count = Counter(r for r, _, _ in planned)
    dup_rows = sorted(r for r, c in row_count.items() if c > 1)
    if dup_rows:
        print(f'!! 重复行号 {dup_rows}（同一目标行被多次填写，后写覆盖先写）')
        for r, d, label in planned:
            if row_count[r] > 1:
                print(f'   {label} -> 行{r} total={d["total"]}')
        print('   多为 decisions 与自动匹配重复配对，或 decisions 来自别的项目')
        return

    total_filled = sum((d['total'] or 0.0) for _, d, _ in planned)
    diff = None if args.expect_total is None else total_filled - args.expect_total
    # 上万条 float 累加必有 1e-8 级误差，按「到分」判定，勿把浮点噪音当真实差异
    tol = 0.0 if diff is None else max(0.01, abs(args.expect_total) * 1e-9)
    aligned = diff is None or abs(diff) <= tol

    print(f'源叶子: item={len(item_map)} desc(item空)={len(desc_map)} '
          f'item0={len(item0)} unit_empty={len(unit_empty)}')
    print(f'计划: match.fill+manual {len(fill)} + green {len(dec["green_fill"]) + len(dec["green_fill_offset"])} '
          f'+ insert {len(dec["green_insert"])} + special {len(dec["special"])}')
    print(f'命中 {fill_ok}, item0兜底 {fill_item0}, 跳过 {fill_skip}')
    print(f'总填 {len(planned)} 行, total 合计 = {total_filled}')
    if diff is not None:
        verdict = '一致' if aligned else '不一致'
        print(f'期望 total = {args.expect_total}, 差额 = {diff:.4f} ({verdict}，容差 {tol:.4f})')

    if skipped_list:
        print(f'\n=== 跳过 {len(skipped_list)} 条 ===')
        for s in skipped_list[:40]:
            print(f'  {s}')
        if len(skipped_list) > 40:
            print(f'  ...另 {len(skipped_list) - 40} 条')

    # 裁决表里的编号大面积落空，通常意味着这份 decisions 不是本项目的
    dec_count = sum(len(dec[k]) for k in DECISION_KEYS)
    if dec_count and len(skipped_list) > max(fill_ok, 5):
        print(f'\n!! 裁决表 {dec_count} 条中大量落空（跳过 {len(skipped_list)} > 命中 {fill_ok}）')
        print('   要么 decisions 不属于本项目，要么源/目标列映射指错列')

    if args.dry_run:
        return

    if not aligned:
        raise SystemExit(f'\n[错误] 差额 {diff:.4f} 超出容差 {tol:.4f}，拒绝写盘。'
                         f'先用 --dry-run 核对 decisions 与列映射')

    bak = args.target.replace('.xlsx', '_backup_before_fill.xlsx')
    shutil.copy2(args.target, bak)
    for row, d, label in planned:
        write_target(ws, row, d, tgt_cols)
        if label.startswith(('green ', 'insert ')):
            for c in range(1, max(tgt_cols.values()) + 2):
                ws.cell(row, c).fill = GREEN
    wb.save(args.target)
    print(f'\n已保存: {args.target}')
    print(f'备份: {bak}')


if __name__ == '__main__':
    main()
