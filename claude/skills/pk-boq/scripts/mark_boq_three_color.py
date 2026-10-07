# -*- coding: utf-8 -*-
"""BOQ 三色差异标记：把新版清单变更以红/黄/绿三色标到基准清单副本。

全程 zip/XML，不启动 Excel。插行、改值、整行上色都走 xlsx_rowops.SheetEditor，
插入行之后的公式行引用由它统一平移（这曾经是只能用 Excel COM 的唯一理由）。

三色语义（整行）：
  绿 = 新增（整行插入到同一 L3 {} 或最接近位置）
  黄 = 修改（改单位/工程量后标黄，不改名称）
  红 = 删除（仅标红不删，等人工删除）

判定规则：
  - 工程量或单位实质变化 -> 修改(黄)
  - 名称/单位同义仅翻译差异 -> 不算偏差，不改（qty 相同即可匹配）
  - 新版独有 -> 新增(绿)
  - 基准独有 -> 删除(红)

Usage:
    # 先 dry-run 看 diff 结果，不写文件
    python mark_boq_three_color.py --base old.xlsx --new new.xlsx \
        --base-cols desc=5,unit=6,qty=7 --new-cols desc=3,unit=4,qty=5 --dry-run

    # 确认后应用（生成三色副本）
    python mark_boq_three_color.py --base old.xlsx --new new.xlsx \
        --base-cols desc=5,unit=6,qty=7 --new-cols desc=3,unit=4,qty=5 -o out.xlsx
"""
import argparse
import json
import re
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

# RGB 十六进制（xlsx 里的写法）。以前是 COM 的 Interior.Color，那是 BGR，
# 所以黄写成 0x00FFFF、红写成 0x0000FF —— 换成 zip/XML 后按 RGB 写。
GREEN = "00FF00"
YELLOW = "FFFF00"
RED = "FF0000"

UNIT_SYNONYMS = {
    ("no", "set"), ("no.", "set"), ("nos", "set"), ("no.", "no"),
    ("no", "no."), ("no.", "nos"), ("no", "nos"),
}

CODE_PAT = re.compile(r"^(tr\d+[a-z]?|is|mh|op|dp)\b")

HEADER_DESC = {"item", "description", "ref", "ref."}


def norm(s):
    if s is None:
        return ""
    s = str(s).strip()
    s = re.sub(r"^\d+\.\d+\s+", "", s)
    s = re.sub(r"\s+", " ", s).lower().strip()
    return s


def tonum(q):
    if q is None:
        return None
    s = str(q).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def norm_l2(s):
    if not s:
        return ""
    s = s.strip("《》").strip().lower()
    s = re.sub(r"\s*\(cont'd\)", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_unit(u):
    return str(u).strip().lower() if u else ""


def same_unit(a, b):
    a, b = norm_unit(a), norm_unit(b)
    if a == b:
        return True
    return (a, b) in UNIT_SYNONYMS or (b, a) in UNIT_SYNONYMS


def code_of(norm_desc):
    m = CODE_PAT.match(norm_desc)
    return m.group(1) if m else ""


def parse_cols(s):
    d = {}
    for part in s.split(","):
        k, _, v = part.partition("=")
        d[k.strip()] = int(v.strip())
    return d


def parse_l2_map(s):
    d = {}
    for part in s.split(","):
        part = part.strip()
        if not part:
            continue
        k, _, v = part.partition("=")
        d[k.strip().lower()] = v.strip().lower()
    return d


def load_rows(path, sheet, cols):
    desc_c = cols.get("desc")
    unit_c = cols.get("unit")
    qty_c = cols.get("qty")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = []
    cur_l2 = None
    for r in range(1, ws.max_row + 1):
        d = ws.cell(row=r, column=desc_c).value if desc_c else None
        if d is None:
            continue
        d = str(d).strip()
        if d.startswith("《"):
            cur_l2 = norm_l2(d)
            continue
        if d.startswith(("{", "【")):
            continue
        if norm(d) in HEADER_DESC:
            continue
        if qty_c is None:
            q = 1
        else:
            q = ws.cell(row=r, column=qty_c).value
            if q is None or str(q).strip() == "":
                continue
        u = str(ws.cell(row=r, column=unit_c).value).strip() if unit_c else ""
        out.append(dict(row=r, desc=d, norm=norm(d), unit=u,
                        qty=q, nq=tonum(q), l2=cur_l2))
    wb.close()
    return out


def match_score(x, y):
    nqx, nqy = x["nq"], y["nq"]
    r = SequenceMatcher(None, x["norm"], y["norm"]).ratio()
    if nqx is not None and nqy is not None and abs(nqx - nqy) < 1e-6:
        return 3.0 + r
    cx, cy = code_of(x["norm"]), code_of(y["norm"])
    if cx and cx == cy:
        return 2.0 + r
    return r if r > 0.7 else 0.0


def align_block(a, b):
    na, nb = len(a), len(b)
    dp = [[0.0] * (nb + 1) for _ in range(na + 1)]
    for i in range(na - 1, -1, -1):
        for j in range(nb - 1, -1, -1):
            s = match_score(a[i], b[j])
            take = dp[i + 1][j + 1] + s if s > 0 else -1e9
            dp[i][j] = max(dp[i + 1][j], dp[i][j + 1], take)
    pairs = []
    i = j = 0
    while i < na and j < nb:
        s = match_score(a[i], b[j])
        if s > 0 and dp[i][j] == dp[i + 1][j + 1] + s:
            pairs.append((i, j))
            i += 1
            j += 1
        elif dp[i][j] == dp[i + 1][j]:
            i += 1
        else:
            j += 1
    return pairs


def diff(base_rows, new_rows, l2_extra, l2_merge_to):
    new_g = defaultdict(list)
    for x in new_rows:
        new_g[x["l2"]].append(x)

    base_g = defaultdict(list)
    orphan = []
    for y in base_rows:
        if y["l2"] in l2_extra:
            orphan.append(y)
        else:
            base_g[y["l2"]].append(y)

    modify = []
    new_items = []
    del_items = []

    for l2, a in new_g.items():
        b = base_g.get(l2, [])
        if l2_merge_to and l2 == l2_merge_to:
            b = b + orphan
            orphan = []
        pairs = align_block(a, b)
        v3_idx_to_cur = {i: b[j]["row"] for i, j in pairs}
        cur_row_to_v3_idx = {b[j]["row"]: i for i, j in pairs}

        last_cur_row = None
        for i, x in enumerate(a):
            if i in v3_idx_to_cur:
                last_cur_row = v3_idx_to_cur[i]
                continue
            anchor = last_cur_row
            if anchor is None:
                for k in range(i + 1, len(a)):
                    if k in v3_idx_to_cur:
                        anchor = v3_idx_to_cur[k]
                        break
            new_items.append(dict(desc=x["desc"], unit=x["unit"], qty=x["qty"],
                                  l2=x["l2"], new_row=x["row"], anchor=anchor))

        for i, j in pairs:
            x, y = a[i], b[j]
            qd = x["nq"] is not None and y["nq"] is not None and abs(x["nq"] - y["nq"]) > 1e-6
            ud = not same_unit(x["unit"], y["unit"])
            if qd or ud:
                modify.append(dict(cur_row=y["row"], cur_desc=y["desc"],
                                   new_unit=x["unit"], new_qty=x["qty"],
                                   change_qty=qd, change_unit=ud, l2=x["l2"]))

        for j, y in enumerate(b):
            if y["row"] not in cur_row_to_v3_idx:
                del_items.append(dict(desc=y["desc"], unit=y["unit"], qty=y["qty"],
                                      l2=y["l2"], cur_row=y["row"]))

    for y in orphan:
        del_items.append(dict(desc=y["desc"], unit=y["unit"], qty=y["qty"],
                              l2=y["l2"], cur_row=y["row"]))

    modify.sort(key=lambda m: m["cur_row"])
    del_items.sort(key=lambda d: d["cur_row"])
    new_items.sort(key=lambda n: n["new_row"])
    return modify, new_items, del_items


def apply_changes(base, modify, new_items, del_items, base_cols, base_sheet, output):
    """标色 + 插行，纯 zip/XML，不启动 Excel。

    行号一律按**原表**给：插入的行号偏移由 SheetEditor 统一算，公式里的行引用
    跟着平移。以前靠 Excel COM 的 Rows.Insert() 做这件事，得从后往前插还要占着
    用户的 Excel。
    """
    import sys as _sys
    _sys.path.insert(0, str(Path(__file__).resolve().parent))
    from xlsx_rowops import SheetEditor

    desc_c = base_cols["desc"]
    unit_c = base_cols.get("unit")
    qty_c = base_cols.get("qty")

    ed = SheetEditor(base, base_sheet)

    for m in modify:
        if m["change_unit"] and unit_c and m["new_unit"]:
            ed.set_cell(m["cur_row"], unit_c, m["new_unit"])
        if m["change_qty"] and qty_c and m["new_qty"] is not None:
            ed.set_cell(m["cur_row"], qty_c, m["new_qty"])
        ed.fill_row(m["cur_row"], YELLOW)

    for d in del_items:
        ed.fill_row(d["cur_row"], RED)

    skipped = [n for n in new_items if n["anchor"] is None]
    for n in new_items:
        if n["anchor"] is None:
            continue
        values = {}
        if n["desc"] is not None:
            values[desc_c] = n["desc"]
        if n["unit"] is not None and unit_c:
            values[unit_c] = n["unit"]
        if n["qty"] is not None and qty_c:
            values[qty_c] = n["qty"]
        ed.insert_row(after=n["anchor"], values=values, fill=GREEN)

    stat = ed.save(output)
    print(f"[apply] 插入 {stat['inserted']} 行，标色 {stat['filled']} 行，"
          f"改值 {stat['edited']} 格，末行 {stat['last_row']}")
    return skipped


def main():
    parser = argparse.ArgumentParser(
        description="BOQ 三色差异标记（绿新增/黄修改/红删除）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""示例:
  python mark_boq_three_color.py --base old.xlsx --new new.xlsx \\
      --base-cols desc=5,unit=6,qty=7 --new-cols desc=3,unit=4,qty=5 --dry-run
  python mark_boq_three_color.py --base old.xlsx --new new.xlsx \\
      --base-cols desc=5,unit=6,qty=7 --new-cols desc=3,unit=4,qty=5 -o out.xlsx
        """,
    )
    parser.add_argument("--base", required=True, help="基准清单（旧版合并 BOQ，干净旧版）")
    parser.add_argument("--new", required=True, help="对比清单（新版设计院清单）")
    parser.add_argument("--base-sheet", default="合并报表", help="基准 sheet 名")
    parser.add_argument("--new-sheet", default="MergeSheet", help="新版 sheet 名")
    parser.add_argument("--base-cols", default="desc=5,unit=6,qty=7",
                        help="基准列映射（1-based），如 desc=5,unit=6,qty=7")
    parser.add_argument("--new-cols", default="desc=3,unit=4,qty=5",
                        help="新版列映射（1-based），如 desc=3,unit=4,qty=5")
    parser.add_argument("-o", "--output", default=None, help="输出副本路径")
    parser.add_argument("--dry-run", action="store_true", help="只 diff 不写文件")
    parser.add_argument("--json", default=None, help="diff 结果 JSON 路径")
    parser.add_argument("--l2-extra", default="",
                        help="基准里需归并到合并 L2 的额外 L2 名（逗号分隔，归一化后）")
    parser.add_argument("--l2-merge-to", default="",
                        help="--l2-extra 归并到的新版 L2 名（归一化后）")
    parser.add_argument("--l2-map", default="",
                        help="L2 名称映射 new=base（逗号分隔），处理 singular/plural 措辞差异")
    args = parser.parse_args()

    base_cols = parse_cols(args.base_cols)
    new_cols = parse_cols(args.new_cols)
    base_rows = load_rows(args.base, args.base_sheet, base_cols)
    new_rows = load_rows(args.new, args.new_sheet, new_cols)
    l2_extra = {x.strip().lower() for x in args.l2_extra.split(",") if x.strip()}
    l2_map = parse_l2_map(args.l2_map)
    for x in new_rows:
        if x["l2"] in l2_map:
            x["l2"] = l2_map[x["l2"]]

    modify, new_items, del_items = diff(base_rows, new_rows, l2_extra, args.l2_merge_to)

    print(f"base={len(base_rows)}  new={len(new_rows)}")
    print(f"修改={len(modify)}  新增={len(new_items)}  删除={len(del_items)}")

    if args.json:
        Path(args.json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.json).write_text(json.dumps(
            {"modify": modify, "new": new_items, "delete": del_items},
            ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(f"JSON: {args.json}")

    if args.dry_run:
        print("\n=== 修改（黄）前 15 条 ===")
        for m in modify[:15]:
            print(f"  r{m['cur_row']} <- {m['cur_desc'][:50]!r} qty={m['new_qty']} unit={m['new_unit']}")
        print("\n=== 新增（绿）前 15 条 ===")
        for n in new_items[:15]:
            print(f"  anchor={n['anchor']} <- {n['desc'][:50]!r} | {n['qty']} {n['unit']}")
        print("\n=== 删除（红）前 15 条 ===")
        for d in del_items[:15]:
            print(f"  r{d['cur_row']} <- {d['desc'][:50]!r} | {d['qty']} {d['unit']}")
        return

    output = args.output or f"{Path(args.base).stem}_三色变更.xlsx"
    skipped = apply_changes(args.base, modify, new_items, del_items,
                            base_cols, args.base_sheet, Path(output))
    if skipped:
        print(f"警告: {len(skipped)} 个新增项 anchor=None（整块无匹配基准项），已跳过需人工定位:")
        for n in skipped:
            print(f"  - {n['desc'][:60]!r}")
    print(f"done: {output}")


if __name__ == "__main__":
    main()
