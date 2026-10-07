#!/usr/bin/env python3
"""
Discover BOQ workbook structure and generate norms config.

Column positions are no longer stored — sync_norms.py auto-detects columns
by header name at runtime. The config only records sheet names and row anchors.

Usage:
    python discover_config.py <workbook.xlsx>
"""

import json
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


COLUMN_PATTERNS = [
    ("bq_code", [
        "bq code", "bq_code", "bqcode",
    ]),
    ("norm_code", [
        "norm code", "定额编号", "norm_code", "quota code",
        "quota_code", "normcode", "norm code2", "norm code 2",
    ]),
    ("norm_name", [
        "norm name", "定额名称", "norm_name", "quota name",
        "quota_name", "normname", "norm name2", "norm name 2",
    ]),
    ("desc", [
        "description", "description of works",
        "item description", "work description",
    ]),
    ("unit", [
        "unit", "单位", "uom", "unit2", "unit 2",
    ]),
    ("qty", [
        "quantity", "qty", "数量", "工程量",
    ]),
    ("rate", [
        "rate", "单价", "unit price", "price",
    ]),
]

DISQUALIFIERS = {
    "norm_code": ["code"],
    "norm_name": ["name", "no"],
    "desc": ["norm name", "定额名称", "norm_name",
             "定额编号", "norm code", "norm_code",
             "name", "no", "code"],
    "bq_code": ["code"],
}


def normalize(s):
    return str(s).strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def match_header(cell_value, role):
    v = normalize(cell_value)
    if not v:
        return False

    for dq in DISQUALIFIERS.get(role, []):
        if normalize(dq) == v:
            return False

    for r, patterns in COLUMN_PATTERNS:
        if r != role:
            continue
        for p in patterns:
            pn = normalize(p)
            if v == pn:
                return True
            if v.startswith(pn):
                return True
            if " " in p and all(w in v for w in pn.split()):
                return True
    return False


def scan_sheet(ws, max_rows=80):
    info = {
        "name": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "headers": [],
        "column_roles": {},
    }

    sample_rows = []
    for row_idx in range(1, min(max_rows + 1, ws.max_row + 1)):
        row_vals = []
        for col_idx in range(1, min(ws.max_column + 1, 50)):
            row_vals.append(ws.cell(row=row_idx, column=col_idx).value)
        sample_rows.append(row_vals)

    header_row = 1
    best_score = 0
    for row_idx in range(min(15, len(sample_rows))):
        row = sample_rows[row_idx]
        score = sum(1 for v in row if v and isinstance(v, str) and len(str(v).strip()) > 1)
        if score > best_score:
            best_score = score
            header_row = row_idx + 1

    info["header_row"] = header_row

    all_text = []
    for col_idx in range(len(sample_rows[header_row - 1])):
        parts = []
        for r_offset in (0, -1, 1):
            r_idx = (header_row - 1) + r_offset
            if 0 <= r_idx < len(sample_rows):
                row_data = sample_rows[r_idx]
                if col_idx < len(row_data):
                    v = row_data[col_idx]
                    if v and str(v).strip():
                        parts.append(str(v).strip())
        all_text.append(" ".join(parts) if parts else "")

    info["headers"] = all_text

    # Detect column roles — store column letter for diagnostic display only
    for role, _ in COLUMN_PATTERNS:
        for idx, h in enumerate(all_text):
            if match_header(h, role):
                col_letter = get_column_letter(idx + 1)
                info["column_roles"][role] = col_letter

    return info


def classify_sheets(workbook_path):
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets_info = []

    for ws in wb.worksheets:
        info = scan_sheet(ws)
        sheets_info.append(info)

    wb.close()

    master_candidates = []
    unique_candidates = []

    for info in sheets_info:
        roles = info["column_roles"]
        has_bq_code = "bq_code" in roles
        has_norm_cols = "norm_code" in roles and "norm_name" in roles

        if info["max_row"] > 5000 and "qty" in roles:
            master_candidates.append(info)
        elif has_bq_code and 100 < info["max_row"] < 10000:
            unique_candidates.append(info)

    master_names = {s["name"] for s in master_candidates}
    unique_candidates = [s for s in unique_candidates if s["name"] not in master_names]

    if not unique_candidates:
        for info in sheets_info:
            roles = info["column_roles"]
            has_desc = "desc" in roles
            if (100 < info["max_row"] < 10000
                    and has_desc
                    and info["name"] not in master_names
                    and info["name"] not in {s["name"] for s in unique_candidates}):
                unique_candidates.append(info)

    result = {
        "file": str(workbook_path),
        "sheets": {s["name"]: s for s in sheets_info},
        "candidates": {
            "master": [s["name"] for s in master_candidates],
            "unique_bq": [s["name"] for s in unique_candidates],
        },
    }
    return result


def build_config(analysis):
    config = {
        "master_sheet": "",
        "master_header_row": 1,
        "master_data_start": 2,
        "unique_bq_sheet": "",
        "unique_bq_header_row": 1,
        "unique_bq_data_start": 2,
    }

    cand = analysis["candidates"]
    sheets = analysis["sheets"]

    if cand["master"]:
        s = sheets[cand["master"][0]]
        config["master_sheet"] = s["name"]
        config["master_header_row"] = s["header_row"]
        config["master_data_start"] = s["header_row"] + 2

    if cand["unique_bq"]:
        s = sheets[cand["unique_bq"][0]]
        config["unique_bq_sheet"] = s["name"]
        config["unique_bq_header_row"] = s["header_row"]
        config["unique_bq_data_start"] = s["header_row"] + 1

    return config


def main():
    if len(sys.argv) < 2:
        print("Usage: python discover_config.py <workbook.xlsx>")
        sys.exit(1)

    wb_path = Path(sys.argv[1])
    if not wb_path.exists():
        print(f"ERROR: File not found: {wb_path}")
        sys.exit(1)

    output_path = wb_path.parent / f"{wb_path.stem}_norms_config.json"

    print(f"Analyzing: {wb_path.name}")
    print("Sheets:")

    analysis = classify_sheets(wb_path)

    for name in sorted(analysis["sheets"].keys()):
        info = analysis["sheets"][name]
        print(f"  [{name}] rows={info['max_row']}, cols={info['max_col']}, "
              f"hdr_row={info['header_row']}, roles={info['column_roles']}")

    config = build_config(analysis)

    master = analysis['candidates']['master']
    unique = analysis['candidates']['unique_bq']

    print(f"\nDetected candidates:")
    print(f"  主清单 (Master BOQ): {master if master else 'NOT FOUND'}")
    print(f"  唯一项表 (Unique items): {unique if unique else 'NOT FOUND'}")

    if len(master) > 1:
        print(f"  WARNING: Multiple master candidates — config uses '{config['master_sheet']}'."
              f"  Edit the JSON if incorrect.")
    if len(unique) > 1:
        print(f"  WARNING: Multiple unique-item candidates — config uses '{config['unique_bq_sheet']}'."
              f"  Edit the JSON if incorrect.")

    output_path.write_text(json.dumps(config, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"\n--- Generated config ---")
    print(json.dumps(config, indent=2, ensure_ascii=False))
    print(f"\nSaved to: {output_path}")
    print("\nColumn positions are auto-detected by sync_norms.py at runtime.")
    print("Expected headers on unique-items sheet: Norm Code2, Norm Name2, Unit2")
    print("Expected headers on master BOQ sheet:   Norm Code, Norm Name, Description, Unit")


if __name__ == "__main__":
    main()
