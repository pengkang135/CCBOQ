#!/usr/bin/env python3
"""
Sync manual norms: unique-items sheet → master BOQ sheet via BQ Code exact match.

Auto-detects columns by header name at runtime — no hardcoded column letters.
Unique-items sheet (pivot table): finds "Norm Code2", "Norm Name2", "Unit2"
  (manual-entry columns at the far right, outside the pivot area).
Master BOQ sheet: finds "Norm Code", "Norm Name", "Description" (first),
  "Unit" (last), "Quantity" (first).

Usage:
    python sync_norms.py <workbook.xlsx> [--dry-run]
"""

import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def normalize(s):
    return str(s).strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def make_bq_code(desc, unit):
    d = str(desc).strip() if desc else ""
    u = str(unit).strip() if unit else ""
    d = re.sub(r'[^\w\s]', '', d)
    return f"{d.lower()}|{u.lower()}"


def _strip_bq_prefix(desc):
    """Replicate T column formula: strip item number prefix from description."""
    d = str(desc).strip() if desc else ""
    if not d:
        return ""
    return re.sub(r'^\s*(?:\d+\.\d+(?:\.\d+)*|-)\s+', '', d)


def _get_combined_header(ws, row, col):
    """Get combined header text from row and neighbour rows (for merged cells)."""
    parts = []
    for r_offset in (-1, 0, 1):
        r = row + r_offset
        if 1 <= r <= ws.max_row:
            val = ws.cell(row=r, column=col).value
            if val:
                parts.append(str(val).strip())
    return " ".join(parts)


def find_column(ws, header_row, patterns, occurrence="first"):
    """Find column by header name patterns.

    patterns: list of strings to match against header text (any match wins).
    occurrence: 'first' returns first match, 'last' returns rightmost match.
    Returns 1-based column index, or None.
    """
    found = None
    for col in range(1, min(ws.max_column + 1, 100)):
        header = _get_combined_header(ws, header_row, col)
        if not header:
            continue
        for p in patterns:
            if normalize(p) in normalize(header):
                if occurrence == "first":
                    return col
                found = col
                break
    return found


def detect_unique_bq_columns(ws, header_row):
    """Auto-detect UniqueBQ columns by header names.

    Prioritises "Norm Code2" / "Norm Name2" / "Unit2" (manual-entry columns
    outside the pivot table). Falls back to rightmost generic match.
    """
    cols = {}

    cols["bq_code"] = find_column(ws, header_row, ["bq code", "bqcode"])

    for role, label_patterns, generic_patterns in [
        ("norm_code", ["norm code2"], ["norm code"]),
        ("norm_name", ["norm name2"], ["norm name"]),
        ("unit",      ["unit2"],      ["unit"]),
    ]:
        col = find_column(ws, header_row, label_patterns, "last")
        if not col:
            col = find_column(ws, header_row, generic_patterns, "last")
        cols[role] = col

    return cols


def detect_zoo_bq_columns(ws, header_row):
    """Auto-detect ZOO BQ columns by header names.

    Description: first occurrence (raw desc with item numbers, col B).
    Unit: last occurrence (leaf unit, col W).
    """
    cols = {}
    cols["desc"]      = find_column(ws, header_row, ["description"], "first")
    cols["unit"]      = find_column(ws, header_row, ["unit"], "last")
    cols["qty"]       = find_column(ws, header_row, ["quantity", "qty"], "first")
    cols["norm_code"] = find_column(ws, header_row, ["norm code"], "first")
    cols["norm_name"] = find_column(ws, header_row, ["norm name"], "first")
    return cols


def load_config(wb_path):
    config_path = wb_path.parent / f"{wb_path.stem}_norms_config.json"
    if config_path.exists():
        return json.loads(config_path.read_text(encoding="utf-8"))
    raise FileNotFoundError(
        f"Config not found: {config_path}\n"
        f"Run discover_config.py first."
    )


def read_unique_bq_norms(wb_path, config):
    """Read UniqueBQ manual norms via openpyxl.

    Returns dict: bq_code -> (norm_code, norm_name).
    """
    sheet = config["unique_bq_sheet"]
    header_row = config.get("unique_bq_header_row", 4)
    data_start = config.get("unique_bq_data_start", 5)

    wb_ro = load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb_ro[sheet]

    cols = detect_unique_bq_columns(ws, header_row)
    print(f"  UniqueBQ columns: {cols}")

    required = ["bq_code", "norm_code"]
    if not all(cols.get(k) for k in required):
        missing = [k for k in required if not cols.get(k)]
        print(f"  ERROR: Could not detect UniqueBQ columns: {missing}")
        wb_ro.close()
        return {}

    c = cols
    involved = [c["bq_code"], c["norm_code"], c.get("norm_name") or c["norm_code"]]
    min_c, max_c = min(involved), max(involved)
    bq_offset  = c["bq_code"] - min_c
    nc_offset  = c["norm_code"] - min_c
    nn_offset  = (c.get("norm_name") or c["norm_code"]) - min_c

    norms = {}
    for row in ws.iter_rows(min_row=data_start, min_col=min_c, max_col=max_c,
                            values_only=True):
        bq_code   = str(row[bq_offset]).strip() if len(row) > bq_offset and row[bq_offset] else ""
        norm_code = str(row[nc_offset]).strip() if len(row) > nc_offset and row[nc_offset] else ""
        norm_name = str(row[nn_offset]).strip() if len(row) > nn_offset and row[nn_offset] else ""
        if bq_code and norm_code:
            norms[bq_code] = (norm_code, norm_name)

    wb_ro.close()
    return norms


def read_zoo_bq_rows(wb_path, config):
    """Read ZOO BQ via openpyxl, compute BQ Code from B+W in Python.

    Returns list: [(excel_row, bq_code), ...]
    """
    sheet = config["master_sheet"]
    header_row = config.get("master_header_row", 3)
    data_start = config.get("master_data_start", 5)

    wb_ro = load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb_ro[sheet]

    cols = detect_zoo_bq_columns(ws, header_row)
    print(f"  ZOO BQ columns: {cols}")

    required = ["desc", "unit"]
    if not all(cols.get(k) for k in required):
        missing = [k for k in required if not cols.get(k)]
        print(f"  ERROR: Could not detect ZOO BQ columns: {missing}")
        wb_ro.close()
        return []

    c = cols
    involved = [c["desc"], c["unit"], c.get("qty") or c["desc"]]
    min_c, max_c = min(involved), max(involved)
    d_offset = c["desc"] - min_c
    u_offset = c["unit"] - min_c
    q_offset = (c.get("qty") or c["desc"]) - min_c

    zoo_rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=data_start, min_col=min_c,
                                           max_col=max_c, values_only=True)):
        excel_row = data_start + idx
        desc = row[d_offset] if len(row) > d_offset else ""
        unit = row[u_offset] if len(row) > u_offset else ""
        qty  = row[q_offset] if len(row) > q_offset else None

        leaf_desc = _strip_bq_prefix(desc)
        if not leaf_desc:
            continue

        bq_code = make_bq_code(leaf_desc, unit)
        if not bq_code or bq_code == "|":
            continue

        zoo_rows.append((excel_row, bq_code))

    wb_ro.close()
    return zoo_rows


def sync(wb_path, dry_run=False):
    config = load_config(wb_path)
    print(f"Workbook: {wb_path.name}")

    # Phase 1: Read
    print("Phase 1: Reading...")
    unique_norms = read_unique_bq_norms(wb_path, config)
    print(f"  UniqueBQ norms: {len(unique_norms)}")

    zoo_rows = read_zoo_bq_rows(wb_path, config)
    print(f"  ZOO BQ rows: {len(zoo_rows)}")

    # Phase 2: Match
    print("Phase 2: Matching...")
    master_ops = []
    for excel_row, bq_code in zoo_rows:
        if bq_code in unique_norms:
            norm_code, norm_name = unique_norms[bq_code]
            master_ops.append((excel_row, norm_code, norm_name))
    print(f"  Matched: {len(master_ops)}")

    # Phase 3: Write
    if dry_run:
        print(f"\n[DRY RUN] Would write {len(master_ops)} norms to ZOO BQ")
        return

    if not master_ops:
        print("\nNo operations to perform.")
        return

    print("Phase 3: Writing...")
    wb = load_workbook(wb_path)
    ws_m = wb[config["master_sheet"]]
    header_row = config.get("master_header_row", 3)
    m_cols = detect_zoo_bq_columns(ws_m, header_row)

    nc_col = m_cols["norm_code"]
    nn_col = m_cols["norm_name"]
    if not nc_col or not nn_col:
        print("  ERROR: Could not detect ZOO BQ norm columns for write")
        wb.close()
        return

    for row, code, name in master_ops:
        ws_m.cell(row=row, column=nc_col, value=code)
        ws_m.cell(row=row, column=nn_col, value=name)

    print(f"  Wrote {len(master_ops)} norms to "
          f"{get_column_letter(nc_col)}/{get_column_letter(nn_col)}")

    output_path = wb_path.parent / f"{wb_path.stem}_synced.xlsx"
    wb.save(str(output_path))
    print(f"\nSaved to: {output_path}")
    wb.close()


def main():
    wb_path = None
    dry_run = "--dry-run" in sys.argv
    for arg in sys.argv[1:]:
        if arg.endswith(".xlsx") or arg.endswith(".xlsm"):
            wb_path = Path(arg)
            break
    if not wb_path:
        print("Usage: python sync_norms.py <workbook.xlsx> [--dry-run]")
        sys.exit(1)
    sync(wb_path, dry_run)


if __name__ == "__main__":
    main()
