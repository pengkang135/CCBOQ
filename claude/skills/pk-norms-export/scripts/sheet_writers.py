"""
Sheet layout and styling for 综合单价 + 人材机.

Everything about how the two sheets look lives here: column positions, headers,
fonts, fills, borders, and the two writers. export_norms.py owns the data.
"""
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from resource_taxonomy import MAIN_CAT2_ORDER, MECH_CAT2_ORDER, MECH_SMALL_CAT3_ORDER

SHEET_RATE = '综合单价'
SHEET_LMM = '人材机'

N_CONSUMPTION = 55
FIXED_COLS = 11
QTY_COL_START = FIXED_COLS + 1          # L
MAX_COL = FIXED_COLS + N_CONSUMPTION    # BN
COL_SEQ, COL_CAT, COL_CODE, COL_NAME, COL_UNIT = 1, 2, 3, 4, 5
COL_AMOUNT, COL_QTY, COL_RATE, COL_LAB, COL_MAT, COL_MECH = 6, 7, 8, 9, 10, 11

EN_HEADERS = ['No.', 'Category', 'Code', 'Description', 'Unit', 'Amount', 'Qty',
              'Rate', 'Labourt', 'Materials', 'Plant']
CN_HEADERS = ['序号', '分类指标', '定额编号', '项目名称', '单位', '合价', '工程量',
              '单价', '人工', '材料', '机械']

LMM_HEADERS = ['编号', '名称', '单位', '数量', '除税单价', '合价', '价格来源']
LMM_COLS = len(LMM_HEADERS)
LMM_QTY, LMM_RATE, LMM_AMOUNT, LMM_SOURCE = 4, 5, 6, 7
CAT1_ORDER = ['人工', '主材', '周转材', '辅材', '机械']


FONT_NAME = '微软雅黑'
FONT_HDR = Font(name=FONT_NAME, size=9, bold=True)
FONT_HDR_EN = Font(name=FONT_NAME, size=8, bold=True)
FONT_CH = Font(name=FONT_NAME, size=9, bold=True, color='FFFFFF')
FONT_SEC = Font(name=FONT_NAME, size=9, bold=True, color='1A1A1A')
FONT_SUB = Font(name=FONT_NAME, size=9, bold=True, color='1A1A1A')
FONT_DATA = Font(name=FONT_NAME, size=9)
FONT_GRAY = Font(name=FONT_NAME, size=9, color='808080')
FONT_TITLE = Font(name=FONT_NAME, size=12, bold=True)
FONT_CAT1 = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
FONT_CAT2 = Font(name=FONT_NAME, size=10, bold=True, color='1F4E79')
FONT_CAT3 = Font(name=FONT_NAME, size=9, bold=True, color='1A1A1A')

FILL_HDR = PatternFill('solid', fgColor='FFD9E1F2')
FILL_CH = PatternFill('solid', fgColor='FF333F4F')
FILL_SEC = PatternFill('solid', fgColor='FFD9E1F2')
FILL_SUB = PatternFill('solid', fgColor='FFFBE5D6')
FILL_CAT2 = PatternFill('solid', fgColor='FFD6E4F0')
# Grey band marking where the 55 consumption columns begin
FILL_QTY_HDR = PatternFill('solid', fgColor=Color(theme=0, tint=-0.15))

_S = Side('thin')
THIN = Border(left=_S, right=_S, top=_S, bottom=_S)
ALIGN_C = Alignment(horizontal='center', vertical='center')
ALIGN_L = Alignment(horizontal='left', vertical='center')
ALIGN_TOP = Alignment(vertical='top')
ALIGN_CT = Alignment(horizontal='center', vertical='top')
ALIGN_HDR = Alignment(horizontal='center', vertical='top', wrap_text=True)
ALIGN_HDR_FIXED = Alignment(horizontal='center', vertical='center', wrap_text=True)


def rate_block_border(ci, top=True, bottom=True):
    """单价 = 人工+材料+机械 reads as one block: no vertical rules inside H:K.
    The two header rows likewise read as one box, so neither draws the rule between them.
    """
    return Border(left=_S if ci not in (COL_LAB, COL_MAT, COL_MECH) else None,
                  right=_S if ci not in (COL_RATE, COL_LAB, COL_MAT) else None,
                  top=_S if top else None, bottom=_S if bottom else None)


BD_DATA = {ci: rate_block_border(ci) for ci in range(1, MAX_COL + 1)}
BD_HDR_TOP = {ci: rate_block_border(ci, bottom=False) for ci in range(1, MAX_COL + 1)}
BD_HDR_BOTTOM = {ci: rate_block_border(ci, top=False) for ci in range(1, MAX_COL + 1)}

NUM_QTY = '#,##0.0000'
NUM_MONEY = '#,##0.00'


def write_lmm(ws, hierarchy, currency, has_qty):
    """人材机 sheet, no merged cells anywhere. Returns {(name, unit): row} for links."""
    ws.cell(1, 1, value='人材机汇总表').font = FONT_TITLE

    ws.cell(2, LMM_COLS - 1, value='币种：').font = FONT_HDR
    ws.cell(2, LMM_COLS - 1).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(2, LMM_COLS, value=currency or None).font = FONT_HDR
    ws.cell(2, LMM_COLS).alignment = ALIGN_L

    for ci, h in enumerate(LMM_HEADERS, 1):
        c = ws.cell(3, ci, value=h)
        c.font = FONT_HDR
        c.fill = FILL_HDR
        c.border = THIN
        c.alignment = ALIGN_C

    def band(row, text, font, fill, level):
        ws.cell(row, 2, value=text).font = font
        for ci in range(1, LMM_COLS + 1):
            ws.cell(row, ci).fill = fill
            ws.cell(row, ci).border = THIN
            ws.cell(row, ci).alignment = ALIGN_L
        ws.row_dimensions[row].outline_level = level

    row, seq = 4, 0
    res_rows = {}
    for cat1 in CAT1_ORDER:
        if cat1 not in hierarchy:
            continue
        band(row, f'【{cat1}】', FONT_CAT1, FILL_CH, 0)
        row += 1

        if cat1 in ('主材', '周转材', '辅材'):
            order = MAIN_CAT2_ORDER
        elif cat1 == '机械':
            order = MECH_CAT2_ORDER
        else:
            order = []
        cat2_items = sorted(hierarchy[cat1].items(),
                            key=lambda x: order.index(x[0]) if x[0] in order else 99)

        for cat2, cat3_d in cat2_items:
            band(row, f'《{cat2}》', FONT_CAT2, FILL_CAT2, 1)
            row += 1
            if cat1 == '机械' and cat2 == '其他中小型机械':
                cat3_items = sorted(cat3_d.items(),
                                    key=lambda x: MECH_SMALL_CAT3_ORDER.index(x[0])
                                    if x[0] in MECH_SMALL_CAT3_ORDER else 99)
            else:
                cat3_items = sorted(cat3_d.items())

            for cat3, items in cat3_items:
                band(row, f'{{{cat3}}}', FONT_CAT3, FILL_SUB, 2)
                row += 1
                for it in items:
                    seq += 1
                    ws.cell(row, 1, value=seq).font = FONT_DATA
                    ws.cell(row, 1).alignment = ALIGN_C
                    ws.cell(row, 2, value=it['name']).font = FONT_DATA
                    ws.cell(row, 3, value=it['unit']).font = FONT_DATA
                    ws.cell(row, 3).alignment = ALIGN_C
                    if has_qty:
                        ws.cell(row, LMM_QTY, value=it['qty'])
                    ws.cell(row, LMM_QTY).font = FONT_DATA
                    ws.cell(row, LMM_QTY).number_format = NUM_MONEY
                    if it['rate'] is not None:
                        ws.cell(row, LMM_RATE, value=round(it['rate'], 4))
                    ws.cell(row, LMM_RATE).font = FONT_DATA
                    ws.cell(row, LMM_RATE).number_format = NUM_MONEY
                    ws.cell(row, LMM_AMOUNT, value=f'=D{row}*E{row}').font = FONT_DATA
                    ws.cell(row, LMM_AMOUNT).number_format = NUM_MONEY
                    ws.cell(row, LMM_SOURCE, value=it['source'] or None).font = FONT_DATA
                    for ci in range(1, LMM_COLS + 1):
                        ws.cell(row, ci).border = THIN
                    ws.row_dimensions[row].outline_level = 3
                    res_rows[(it['name'], it['unit'])] = row
                    row += 1

    for cl, w in {'A': 8, 'B': 40, 'C': 10, 'D': 14, 'E': 14, 'F': 14, 'G': 22}.items():
        ws.column_dimensions[cl].width = w
    ws.freeze_panes = 'A4'
    ws.sheet_properties.outlinePr.summaryBelow = False
    return res_rows


def write_rate_sheet(ws, tree, qty_map, res_rows, has_qty):
    """综合单价 sheet with SUMPRODUCT rates and 人材机 links."""
    def style_row(row, font, fill=None):
        for ci in range(1, MAX_COL + 1):
            c = ws.cell(row, ci)
            c.font = font
            if fill:
                c.fill = fill
            c.border = BD_DATA[ci]
            # 单位 onwards is numeric or short: centred, hugging the top
            c.alignment = ALIGN_CT if ci >= COL_UNIT or ci == COL_SEQ else ALIGN_TOP

    for row, headers, border in ((1, EN_HEADERS, BD_HDR_TOP), (2, CN_HEADERS, BD_HDR_BOTTOM)):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row, ci, value=h)
            c.font = FONT_HDR
            c.border = border[ci]
            c.alignment = ALIGN_HDR if ci >= COL_UNIT else ALIGN_HDR_FIXED
        for ci in range(QTY_COL_START, MAX_COL + 1):
            n = ci - FIXED_COLS
            c = ws.cell(row, ci, value=f'Qty{n}' if row == 1 else f'消耗量{n}')
            c.font = FONT_HDR_EN
            c.fill = FILL_QTY_HDR
            c.border = border[ci]
            c.alignment = ALIGN_HDR

    row, seq = 3, 1
    for ch_name, ch_d in tree.items():
        ws.cell(row, COL_CODE, value=ch_name.split()[0] if ch_name else None)
        ws.cell(row, COL_NAME, value=f'【{ch_name}】')
        style_row(row, FONT_CH, FILL_CH)
        row += 1

        for sec_name, subs in ch_d['sections'].items():
            ws.cell(row, COL_NAME, value=f'《{sec_name}》' if sec_name else None)
            style_row(row, FONT_SEC, FILL_SEC)
            ws.row_dimensions[row].outline_level = 1
            row += 1

            for sub_name, sub_d in subs.items():
                resources = sub_d['resources']
                cols = {'labour': [], 'materials': [], 'mech': []}
                for ri, res in enumerate(resources):
                    cols[res['cls']].append(QTY_COL_START + ri)

                ws.cell(row, COL_NAME, value=f'{{{sub_name}}}' if sub_name else None)
                for ri, res in enumerate(resources):
                    ws.cell(row, QTY_COL_START + ri, value=res['label'])
                style_row(row, FONT_SUB, FILL_SUB)
                ws.row_dimensions[row].outline_level = 2
                row += 1

                price_row = row
                ws.cell(row, COL_NAME, value='参考单价 ->').font = FONT_GRAY
                for ri, res in enumerate(resources):
                    c = ws.cell(row, QTY_COL_START + ri)
                    lmm_row = res_rows.get((res['name'], res['unit']))
                    if lmm_row:
                        c.value = f"='{SHEET_LMM}'!{get_column_letter(LMM_RATE)}{lmm_row}"
                    c.number_format = NUM_MONEY
                style_row(row, FONT_GRAY)
                ws.row_dimensions[row].outline_level = 2
                row += 1

                for item in sub_d['items']:
                    ws.cell(row, COL_SEQ, value=seq).alignment = ALIGN_C
                    ws.cell(row, COL_SEQ).number_format = '0'
                    ws.cell(row, COL_CAT, value=sub_name.split(' ', 1)[-1] if sub_name else None)
                    ws.cell(row, COL_CODE, value=item['code']).number_format = '@'
                    ws.cell(row, COL_NAME, value=item['name'])
                    ws.cell(row, COL_UNIT, value=item['units'])

                    ws.cell(row, COL_AMOUNT, value=f'=G{row}*H{row}').number_format = NUM_MONEY
                    if has_qty:
                        q = qty_map.get(item['code'])
                        if q:
                            ws.cell(row, COL_QTY, value=round(q, 4))
                    ws.cell(row, COL_QTY).number_format = NUM_QTY
                    ws.cell(row, COL_RATE, value=f'=I{row}+J{row}+K{row}').number_format = NUM_MONEY

                    for cls, col in (('labour', COL_LAB), ('materials', COL_MAT), ('mech', COL_MECH)):
                        cc = cols[cls]
                        if not cc:
                            continue
                        a, b = get_column_letter(cc[0]), get_column_letter(cc[-1])
                        ws.cell(row, col,
                                value=f'=IFERROR(SUMPRODUCT({a}{price_row}:{b}{price_row},{a}{row}:{b}{row}),0)')
                        ws.cell(row, col).number_format = NUM_MONEY

                    amounts = {r['label']: r['amount'] for r in item['resources']}
                    for ri, res in enumerate(resources):
                        val = amounts.get(res['label'])
                        if val:
                            c = ws.cell(row, QTY_COL_START + ri, value=val)
                            c.number_format = NUM_QTY
                    style_row(row, FONT_DATA)
                    ws.row_dimensions[row].outline_level = 3
                    row += 1
                    seq += 1
                row += 1
            row += 1
        row += 1

    # 单位 (E) stays narrow; 人工/材料/机械 (I-K) keep Excel's default width
    widths = {'A': 6, 'B': 14, 'C': 18.22, 'D': 44.22, 'E': 5.11,
              'F': 12, 'G': 12, 'H': 12}
    for cl, w in widths.items():
        ws.column_dimensions[cl].width = w
    for ci in range(QTY_COL_START, MAX_COL + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 13
    ws.freeze_panes = 'E3'
    ws.sheet_properties.outlinePr.summaryBelow = False
    return seq - 1
